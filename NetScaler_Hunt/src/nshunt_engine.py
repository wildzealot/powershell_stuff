#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
nshunt_engine.py - hunt engine used by Invoke-NetScalerHunt.ps1 (it can also be run on its own).

Opens Citrix NetScaler (ADC / Gateway) disk images READ-ONLY with dissect, exports /var/log,
converts every exported log to .xlsx, runs the hunts, and writes a results workbook per appliance
plus a Word findings summary. Nothing is ever written inside the evidence folder.

Supported images: VMware VMDK (snapshot chains are followed), raw dd images (.dd, .raw, .img),
split raw images (.001, .002, ...) and EnCase E01 (.E01, .E02, ...). --path is either a folder
(searched three levels deep for all of these) or one image file (any extension is read as raw).

  python -B nshunt_engine.py --path <folder or image file> --out <output folder> [options]

Examples (the output folder must not be inside the evidence folder):
  Case folder, every supported image in it and up to three levels below:
    python -B nshunt_engine.py --path D:\Evidence\Case01 --out C:\NSHunt\Case01
  One VMware VM (give the newest snapshot disk to open the current state, or the base disk for the original):
    python -B nshunt_engine.py --path D:\Evidence\Case01\ns-vpx01\ns-vpx01-000002.vmdk --out C:\NSHunt\ns-vpx01
  Raw dd image (.raw and .img work the same way):
    python -B nshunt_engine.py --path D:\Evidence\Case01\ns01.dd --out C:\NSHunt\ns01
  Split raw image (give the first part; .002, .003 ... are picked up from the same folder):
    python -B nshunt_engine.py --path D:\Evidence\Case01\ns02.001 --out C:\NSHunt\ns02
  EnCase image (give the .E01; .E02, .E03 ... are picked up from the same folder):
    python -B nshunt_engine.py --path D:\Evidence\Case01\ns03.E01 --out C:\NSHunt\ns03
  Linux or macOS, with extra attacker IPs for this client:
    python3 -B nshunt_engine.py --path /cases/case01/ns01.dd --out ~/nshunt/ns01 --ioc-ip 203.0.113.10 198.51.100.7

Options:
  --since YYYY-MM-DD    start of the review window (default: 30 days before the newest log of each image)
  --ioc-ip IP [IP ...]  extra attacker IPs (IPv4 or IPv6) to search every log for; they are added to the built-in
                        list (BUILTIN_IOC_IPS below: the tool's own case plus public reporting), which is always searched
  --foxit PATH          Fox-IT iocitrix.py to run with the compatibility fixes (omit to skip)
  --skip-log-xlsx       do not convert every log to .xlsx (the hunts still run)
  --include-base        also run the full hunt on the base (pre-snapshot) disk of each chain
All times are UTC. Syslog lines carry no year or zone; the year comes from the file's modified time and
the appliance clock is assumed to be UTC (the engine checks this against ns.log GMT stamps).
"""
import argparse
import base64
import binascii
import collections
import contextlib
import csv
import datetime as dt
import fnmatch
import gzip
import hashlib
import importlib.util
import ipaddress
import json
import logging
import os
import platform
import re
import stat
import struct
import sys
import time
import traceback
import urllib.parse
from pathlib import Path

ENGINE_VERSION = "1.3.1"
UTC = dt.timezone.utc
XL_MAX_ROWS = 1048575      # data rows per sheet (one header row on top)
XL_MAX_CELL = 32000        # Excel allows 32,767 characters per cell

# ------------------------------------------------------------------ known indicators (edit here)
# Every run searches all logs for the attacker IPs below (hunt H); --ioc-ip adds more and never replaces them.
# A hit is a lead to check, not proof: read the log line before calling it attacker activity.
#
# 1) The October 2026 intrusion this tool was built on (CVE-2026-88771 log poisoning, then a webshell, tunnels and
#    config theft). Level: confirmed = carried out attacker actions; likely = requested the webshell minutes before a
#    config theft; suspected = only probed this implant's webshell paths (could be a scanner that knows them).
CASE_IOC_IPS = [
    ("213.209.159.55", "confirmed", "sent the poisoned logins, hosted the payload script that Call Home ran, used the "
                                    "webshell (2026-10; also reported by Beazley and Rapid7)"),
    ("124.198.132.44", "confirmed", "used the webshell as a tunnel into the internal network; SSH server of the reverse "
                                    "SSH tunnels (2026-10)"),
    ("162.33.177.119", "confirmed", "tried 118 times to download a config archive that an injected command was meant to "
                                    "create (2026-10)"),
    ("152.233.50.151", "likely", "requested the webshell minutes before a config theft (2026-10)"),
    ("146.70.211.157", "suspected", "probed this implant's webshell paths (2026-10); Huntback.io decoys saw it send "
                                    "CVE-2026-88771 injection checks"),
    ("23.162.8.173", "suspected", "probed this implant's webshell paths (2026-10); Huntback.io decoys saw it send "
                                  "CVE-2026-88771 injection checks"),
    ("37.19.221.148", "suspected", "probed this implant's webshell paths (2026-10)"),
    ("104.36.50.29", "suspected", "probed this implant's webshell paths (2026-10)"),
    ("23.234.94.126", "suspected", "probed this implant's webshell paths (2026-10)"),
    ("68.235.46.82", "suspected", "probed this implant's webshell paths (2026-10)"),
]
# 2) Public reporting on NetScaler exploitation up to REPORTED_AS_OF: Rapid7, Zscaler, watchTowr, Unit 42, LevelBlue,
#    eSentire, Mandiant, GreyNoise, Arctic Wolf, Sygnia, Truesec, Beazley, TENEX, SOCRadar, Lupovis, Corelight, IFIN, and
#    the Poppelgaard checker (which relays Gotham Technology Group IR and Huntback.io decoy data). Only C2, payload,
#    staging and exfiltration servers (reported-infra), webshell users and exploit sources named by vendor or IR
#    reports (reported-exploit), and IPs in a vendor's IOC table with no stated role (reported-ioc) are included;
#    mass-scan feeds and Tor exits are left out, and so are Cloudflare WARP addresses except the four that Unit 42 tied
#    to this actor's webshell and GetUserName requests (ordinary WARP users share those too, so a hit on one of them
#    is weak on its own). One line per IP: ip | level | what it was seen doing (sources).
REPORTED_AS_OF = "2026-10-19"
_REPORTED_IOC_IPS = """
38.180.81.157 | reported-infra | Platypus C2 node sharing the cluster's TLS certificate (TENEX, Poppelgaard checker)
45.141.21.130 | reported-infra | reverse-shell C2 (LevelBlue, eSentire, Arctic Wolf, TENEX, Poppelgaard checker)
45.143.130.195 | reported-infra | C2 and payload server of the NetScaler C2 mass-exploitation framework (8899/tcp HTTP, 53/udp DNS) (SOCRadar, Poppelgaard checker)
77.83.199.39 | reported-infra | potential C2 seen by Truesec (Unit 42, eSentire, Sygnia, Truesec, IFIN, Poppelgaard checker)
89.44.80.7 | reported-infra | callback host (:65456 /exec-ok and id output; nc shell to :58963) (Arctic Wolf, Poppelgaard checker)
95.133.231.109 | reported-infra | Platypus C2 node sharing the cluster's TLS certificate (TENEX, Poppelgaard checker)
104.200.67.56 | reported-infra | Platypus C2 node sharing the cluster's TLS certificate (TENEX, Poppelgaard checker)
104.248.244.66 | reported-infra | potential C2 seen by Truesec (Unit 42, Truesec, IFIN, Poppelgaard checker)
130.94.20.222 | reported-infra | payload delivery (Rapid7 2026-10-19, watchTowr, TENEX, Poppelgaard checker)
130.94.42.226 | reported-infra | callback host (:18805) (Arctic Wolf, Poppelgaard checker)
139.180.152.138 | reported-infra | potential C2 seen by Truesec (Unit 42, eSentire, Truesec, IFIN, Poppelgaard checker)
176.65.148.54 | reported-infra | Sliver C2 endpoint (:80) and ns_helper staging (:8080) (watchTowr, GreyNoise, Poppelgaard checker)
192.159.99.139 | reported-infra | .slap C2, relay, reverse-agent and loot-upload infrastructure (Rapid7 2026-10-19)
194.26.29.88 | reported-infra | Host of the reverse shell documented by Corelight (Corelight, Poppelgaard checker)
195.123.233.245 | reported-infra | Primary Platypus C2 node (443); entretiensol.com and white-guard.pro resolve here (TENEX, Poppelgaard checker)
199.233.217.13 | reported-infra | payload delivery (Rapid7 2026-10-19, TENEX, Poppelgaard checker)
209.74.64.192 | reported-infra | C2 used by nsvmd (FreeBSD Go proxy in /flash/.nsvm/) (Rapid7 2026-10-19)
45.92.1.17 | reported-infra | stage/loot infrastructure on port 30444 (Rapid7 2026-10-19)
64.94.85.67 | reported-infra | payload delivery (Rapid7 2026-10-19, watchTowr, LevelBlue, eSentire, Arctic Wolf, Sygnia, TENEX, GreyNoise, Poppelgaard checker)
81.94.239.8 | reported-infra | exfiltration server (port 8877) for a payload that stole ns.conf and the SSL keys (field report, 6 Oct) (Poppelgaard checker)
138.199.200.90 | reported-infra | destination for data exfiltrated through log poisoning (Lupovis) (Lupovis, Poppelgaard checker)
158.94.211.205 | reported-infra | exfiltration receiver (:8080) (Beazley, Poppelgaard checker)
23.27.143.20 | reported-infra | exploit source and payload host (main.py) (LevelBlue, eSentire, Arctic Wolf, TENEX, Poppelgaard checker)
31.56.197.72 | reported-infra | payload delivery (Rapid7 2026-10-19, LevelBlue, eSentire, Arctic Wolf, TENEX, Poppelgaard checker)
46.151.182.18 | reported-infra | payload delivery (Rapid7 2026-10-19)
46.151.182.131 | reported-infra | Bot-download host (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
62.133.62.80 | reported-infra | payload delivery (Rapid7 2026-10-19, LevelBlue, eSentire, Arctic Wolf, TENEX, Poppelgaard checker)
68.178.160.183 | reported-infra | payload host (:8888 /test, :8899 /test111) (eSentire, Arctic Wolf, Poppelgaard checker)
72.5.65.111 | reported-infra | tied to the 8 Oct v5v.in injection wave (payload host or wave source; reports differ) (Poppelgaard checker)
77.247.126.239 | reported-infra | Backing host of the pinggy/serveo tunnel droppers (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
132.243.166.140 | reported-infra | payload delivery (Rapid7 2026-10-19)
134.175.71.50 | reported-infra | payload host (:4123) (Arctic Wolf, Poppelgaard checker)
143.198.7.94 | reported-infra | CVE-2026-88772 scanning and staging infrastructure (Mandiant, Beazley, Poppelgaard checker)
153.75.82.220 | reported-infra | payload host (/download/x.sh) (eSentire, Arctic Wolf, GreyNoise, Poppelgaard checker)
173.231.39.244 | reported-infra | payload delivery (Rapid7 2026-10-19, watchTowr, GreyNoise)
177.4.12.11 | reported-infra | payload host (:8080 /s?t=a/b/c) (Arctic Wolf, Poppelgaard checker)
179.60.146.241 | reported-infra | served 1.sh (Rapid7 2026-10-19)
209.141.43.95 | reported-infra | payload host for kw2-freebsd-x64-s11 (Rust FreeBSD implant); also seen in StyleSmuggler activity (Zscaler 2026-10-16)
45.59.125.187 | reported-exploit | PHP webshell source (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
66.135.19.18 | reported-exploit | requested the nsgclient18.deb / nsgser18.deb webshells (4-8 Sep) (Unit 42, Poppelgaard checker)
104.28.215.136 | reported-exploit | Cloudflare WARP address requesting .deb webshells (9-11 Sep) (Unit 42)
104.28.215.137 | reported-exploit | Cloudflare WARP address requesting .deb webshells (9-11 Sep) (Unit 42)
104.28.247.136 | reported-exploit | Cloudflare WARP address requesting .deb webshells (9-11 Sep) (Unit 42)
104.28.247.137 | reported-exploit | Cloudflare WARP address, continuous GetUserName requests (10-27 Sep) (Unit 42)
104.248.74.206 | reported-exploit | requested the nsgclient18.deb / nsgser18.deb webshells (4-8 Sep) (Unit 42, Poppelgaard checker)
137.184.91.207 | reported-exploit | requested the same .deb webshells (7 Sep) (Unit 42, Poppelgaard checker)
142.93.85.227 | reported-exploit | requested the nsgclient18.deb / nsgser18.deb webshells (4-8 Sep) (Unit 42, Poppelgaard checker)
162.33.178.9 | reported-exploit | requested the nsgbuild.deb / nsgsupport.deb webshells (14 Sep) (Unit 42, Poppelgaard checker)
167.99.111.203 | reported-exploit | requested the nsgclient18.deb / nsgser18.deb webshells (4-8 Sep) (Unit 42, Poppelgaard checker)
193.149.176.207 | reported-exploit | requested nsgbuild.deb daily (15-24 Sep) and GetUserName (Unit 42, Poppelgaard checker)
5.83.144.60 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Poppelgaard checker)
5.188.206.226 | reported-exploit | exploitation source (eSentire, Poppelgaard checker)
23.234.74.48 | reported-exploit | CVE-2026-88771 exploitation source (watchTowr honeypots); Huntback.io decoys saw it prove exploitation out of band (watchTowr, GreyNoise, Poppelgaard checker)
32.186.46.211 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
34.90.151.231 | reported-exploit | reconnaissance and webshell delivery (eSentire, Poppelgaard checker)
37.19.221.171 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, GreyNoise)
44.252.255.141 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, GreyNoise)
45.12.239.191 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
45.61.144.161 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
45.76.34.141 | reported-exploit | high confidence: seen in malicious AAA events (Sygnia, Poppelgaard checker)
45.143.167.96 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, GreyNoise)
45.153.212.101 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
45.225.135.18 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
45.249.89.172 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, GreyNoise)
46.33.33.95 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
46.233.40.20 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
46.236.167.143 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
47.230.224.154 | reported-exploit | source of NetScaler exploitation attempts (LevelBlue, Poppelgaard checker)
49.150.197.119 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
51.158.203.95 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Beazley, Previdian, Poppelgaard checker)
52.38.11.186 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
52.39.16.162 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
54.70.59.128 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, GreyNoise)
64.227.181.23 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Poppelgaard checker)
66.42.100.63 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Poppelgaard checker)
70.172.58.168 | reported-exploit | source of NetScaler exploitation attempts (LevelBlue, Poppelgaard checker)
78.47.24.217 | reported-exploit | fingerprinted 100+ systems (21-22 Aug); part of the 21 Sep webshell drop (Unit 42, Poppelgaard checker)
78.128.113.10 | reported-exploit | exploitation-attempt source (GreyNoise, Marius Sandbu and Lupovis lists, 28-29 Sep) (Poppelgaard checker)
78.128.114.22 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
78.135.96.136 | reported-exploit | Associated exploit source (IFIN compiled observables) (IFIN, Poppelgaard checker)
78.190.152.64 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
80.240.22.229 | reported-exploit | Associated exploit source (IFIN compiled observables) (IFIN, Poppelgaard checker)
81.10.79.197 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
82.167.14.7 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16, eSentire, Beazley, GreyNoise, Poppelgaard checker)
87.224.84.82 | reported-exploit | configuration-staging attempt (LevelBlue, GreyNoise, Poppelgaard checker)
89.36.231.206 | reported-exploit | Associated exploit source (IFIN compiled observables) (IFIN, Poppelgaard checker)
91.195.240.123 | reported-exploit | Associated exploit source (IFIN compiled observables) (IFIN, Poppelgaard checker)
92.118.204.229 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, LevelBlue, eSentire, GreyNoise, Poppelgaard checker)
95.105.74.172 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
100.38.125.66 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
102.213.132.123 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
102.222.146.45 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
103.83.146.100 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
103.181.212.245 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
103.197.249.161 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
104.168.34.24 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
104.207.32.77 | reported-exploit | exploit delivery (Beazley, Poppelgaard checker)
104.207.46.202 | reported-exploit | exploit delivery (Beazley, Poppelgaard checker)
104.207.47.54 | reported-exploit | exploit delivery (Beazley, Poppelgaard checker)
104.234.140.124 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
104.234.140.132 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
109.71.252.97 | reported-exploit | ns.conf-theft source (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
111.92.128.89 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
114.145.49.215 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
120.244.160.182 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
125.62.88.26 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
130.12.182.7 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16, Poppelgaard checker)
130.94.106.141 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, GreyNoise)
137.220.53.135 | reported-exploit | CVE-2026-88771 exploitation attempts (GreyNoise, 28-29 Sep); Huntback.io decoys saw it upload ns.conf (GreyNoise, Poppelgaard checker)
138.28.234.38 | reported-exploit | exploitation with attempted DNS exfiltration (Lupovis) (Beazley, Lupovis, Poppelgaard checker)
138.68.21.29 | reported-exploit | medium confidence: temporal and service-path correlation (Sygnia, Poppelgaard checker)
138.199.60.5 | reported-exploit | sent CVE-2026-88779 crash payloads to /saml/login and /cgi/samlauth on 5 Oct (Gotham Technology Group) (Poppelgaard checker)
138.199.60.22 | reported-exploit | sent CVE-2026-88771 injection checks to Huntback.io decoys (4 Oct) (Poppelgaard checker)
138.199.60.36 | reported-exploit | sent CVE-2026-88771 injection checks to Huntback.io decoys (4 Oct) (Poppelgaard checker)
139.59.86.242 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Poppelgaard checker)
141.98.212.82 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
142.93.205.229 | reported-exploit | CVE-2026-88771 exploitation source (watchTowr honeypots); polled for the c88771.json success marker (watchTowr, GreyNoise, Poppelgaard checker)
144.172.108.78 | reported-exploit | exploitation source (eSentire, Poppelgaard checker)
146.70.199.170 | reported-exploit | sent CVE-2026-88771 injection checks to Huntback.io decoys (4 Oct) (Poppelgaard checker)
149.28.29.221 | reported-exploit | exploitation source (eSentire, IFIN, Poppelgaard checker)
149.104.78.141 | reported-exploit | exploitation source (eSentire, GreyNoise, Beazley, Poppelgaard checker)
149.104.78.208 | reported-exploit | earliest exploitation attempt Rapid7 saw (20 Sep): tar of /flash/nsconfig to a web path (Rapid7, Poppelgaard checker)
154.183.158.198 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
154.217.251.226 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Beazley, GreyNoise, Poppelgaard checker)
154.241.254.9 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
157.254.167.12 | reported-exploit | CVE-2026-88772 exploitation and basic webshell backdoor (Mandiant, Beazley, Poppelgaard checker)
158.94.209.12 | reported-exploit | exploitation source (eSentire, Poppelgaard checker)
159.26.112.64 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
159.65.104.231 | reported-exploit | CVE-2026-88771 exploitation attempts (GreyNoise, 28-29 Sep); polled for the c88771.json success marker (GreyNoise, Poppelgaard checker)
159.203.33.46 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Poppelgaard checker)
162.243.36.88 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, LevelBlue, Poppelgaard checker)
165.22.100.102 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
165.22.104.177 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
165.227.201.112 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, GreyNoise)
165.227.228.21 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Poppelgaard checker)
170.64.143.206 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Poppelgaard checker)
170.64.176.26 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Sygnia, Poppelgaard checker)
172.247.44.85 | reported-exploit | CVE-2026-88771 exploitation source (watchTowr honeypots); Huntback.io decoys saw it prove exploitation out of band (watchTowr, GreyNoise, Poppelgaard checker)
173.40.135.209 | reported-exploit | source of NetScaler exploitation attempts (LevelBlue, Poppelgaard checker)
174.60.63.37 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
177.52.203.23 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
183.63.88.227 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
185.100.87.166 | reported-exploit | Tunnel-dropper source (pinggy/serveo fronted droppers) (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
185.121.170.60 | reported-exploit | ns.conf-theft source (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
185.135.77.63 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
185.156.46.162 | reported-exploit | exploitation source (eSentire, GreyNoise, Poppelgaard checker)
185.231.33.46 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Poppelgaard checker)
185.243.218.225 | reported-exploit | ns.conf-theft source (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
185.244.213.112 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, Beazley, Poppelgaard checker)
190.117.56.56 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
194.54.83.22 | reported-exploit | Out-of-band proving source (dnshook.site callbacks) (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
195.62.37.119 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
198.13.159.233 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, GreyNoise)
204.8.96.74 | reported-exploit | ns.conf-theft source (Huntback.io decoys, 7 Oct) (Poppelgaard checker)
206.189.107.84 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
207.48.184.10 | reported-exploit | source IP of the staged CVE-2026-88771 chain (INDEX: User-Agent staging, pitboss trigger) (Zscaler 2026-10-16)
209.250.236.77 | reported-exploit | medium confidence: time-correlated with exploit-style activity (Sygnia, Poppelgaard checker)
212.86.125.40 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr)
216.203.21.233 | reported-exploit | CVE-2026-88771 exploitation source seen on watchTowr's honeypot sensors (watchTowr, eSentire, GreyNoise, Poppelgaard checker)
23.234.80.246 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
23.234.83.194 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
23.234.109.28 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
31.56.197.137 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
64.176.71.42 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
66.173.222.26 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
79.133.42.141 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
85.11.187.35 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
91.199.163.55 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
103.214.20.54 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
109.136.126.142 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
135.136.98.176 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
139.162.75.170 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
139.162.83.159 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
143.244.44.177 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
146.70.199.53 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
159.223.233.184 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
167.88.172.6 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
194.127.166.126 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
207.148.105.57 | reported-exploit | in a Gotham Technology Group IR group (2 Oct) whose members did recon, webshell polling, payload or callback work (Poppelgaard checker)
45.61.136.143 | reported-ioc | in Unit 42's network IOC table with no stated role; BL Networks, the provider of several other reported IPs (Unit 42, Poppelgaard checker)
66.227.183.84 | reported-ioc | in Unit 42's network IOC table with no stated role; Charter Communications network (Unit 42, Poppelgaard checker)
216.245.184.164 | reported-ioc | in Unit 42's network IOC table with no stated role; BL Networks, the provider of several other reported IPs (Unit 42, Poppelgaard checker)
"""
REPORTED_IOC_IPS = [tuple(x.strip() for x in line.split("|", 2)) for line in _REPORTED_IOC_IPS.strip().splitlines()]
BUILTIN_IOC_IPS = CASE_IOC_IPS + [x for x in REPORTED_IOC_IPS if x[0] not in {c[0] for c in CASE_IOC_IPS}]
SUPPLIED_IOC = ("supplied", "given with -IocIp / --ioc-ip for this run")

KNOWN_IOC_PATHS = [
    ("/flash/nsconfig/.slap", "Hidden implant folder: agent.pl, bridge.pl, boot.sh and copies of the PHP webshell; Rapid7 "
                              "(2026-10-19) also lists reverse-agent.sh, vcscan-agent.sh and vc here"),
    ("/flash/nsconfig/.relay", "Relay component folder (reverse SSH tunnel and SOCKS proxy, guard.sh); Rapid7 (2026-10-19) "
                               "also lists run.sh and target_agent.py here"),
    ("/flash/.nsvm", "nsvmd FreeBSD Go proxy folder (Rapid7, 2026-10-19): nsvmd, watchdog.sh, guard.sh, restart.sh, "
                     ".agent.log, .wd.log; nsvmd talks to v5v.in / 209.74.64.192"),
    ("/var/tmp/.s2loot", "Loot staging folder of the same implant (Rapid7, 2026-10-19); this tool's own case had the log "
                         "file /var/tmp/.s2loot.log instead"),
    ("/var/nstmp/.relay", "Relay component folder (alternate location)"),
    ("/var/tmp/.ux", "Tunnel agent folder: agent.log, slapshot.py, whipd.py and pid files"),
    ("/var/tmp/.s2loot.log", "Log of nsconfig archives made by the implant (archive=nsconfig ... valid=yes)"),
    ("/var/tmp/.relay-ssh.log", "Relay SSH client log (names the SSH server)"),
    ("/var/tmp/.relay-socks.log", "Relay SOCKS proxy log"),
    ("/var/tmp/.slap-diag.txt", "Implant diagnostics output"),
    ("/var/tmp/.slap-watch.pid", "Implant watchdog pid file"),
    ("/var/tmp/.slap-watch.log", "Implant watchdog log"),
    ("/var/tmp/.slap-agent.log", "Implant agent log"),
    ("/var/tmp/.slap-httpd-test.log", "Implant web test log"),
    ("/var/netscaler/logon/LogonPoint/custom/.ctxs.receiver", "Webshell copy, served at /logon/LogonPoint/custom/.ctxs.receiver"),
    ("/var/netscaler/logon/LogonPoint/custom/.slap.receiver", "Webshell copy, served at /logon/LogonPoint/custom/.slap.receiver"),
    ("/var/netscaler/logon/LogonPoint/custom/receiver.deb", "Webshell copy, served at /logon/LogonPoint/custom/receiver.deb"),
    ("/var/cron/tabs/nobody", "Crontab of the nobody user (the web server account)"),
    ("/flash/nsconfig/rc.netscaler.ptbak", "Copy of rc.netscaler kept when rc.netscaler was rewritten (holds the earlier version)"),
]
# Not listed above on purpose: /flash/nsconfig/rc.netscaler and /var/cron/tabs/root. Administrators use both for
# ordinary settings (for example nsapimgr_wr.sh tuning), so their presence is not an indicator. Their lines are
# reviewed on the Hooks sheet instead, and only flagged lines reach the Word summary.
KNOWN_IOC_RX = re.compile(
    r"(^|/)\.(slap|relay|ux)(/|$)|^/flash/\.nsvm(/|$)"
    r"|^/var/tmp/\.(slap-[^/]+|relay-(ssh|socks)\.log)$|^/var/tmp/\.s2loot(\.log$|/|$)"
    r"|/LogonPoint/custom/(\.ctxs\.receiver|\.slap\.receiver|receiver\.deb)$"
    r"|^/var/cron/tabs/nobody$|^/flash/nsconfig/rc\.netscaler\.ptbak$"
)
# Text logs written by the implant. They are exported and converted like /var/log. Implant code is never exported.
IMPLANT_LOG_EXPORTS = [
    "/var/tmp/.ux/agent.log", "/var/tmp/.ux/whip.log", "/var/tmp/.relay-ssh.log", "/var/tmp/.relay-socks.log",
    "/var/tmp/.s2loot.log", "/var/tmp/.slap-diag.txt", "/var/tmp/.slap-agent.log", "/var/tmp/.slap-watch.log",
    "/var/tmp/.slap-httpd-test.log", "/flash/.nsvm/.agent.log", "/flash/.nsvm/.wd.log",
]
IMPLANT_STRINGS = (r"\.slap|/\.ux/|slapshot|whipd|receiver\.deb|ctxs\.receiver|slap\.receiver|CsrfToken|\.relay|guard\.sh"
                   r"|s2loot|nsvmd|\.nsvm|v5v\.in|vcscan|reverse-agent|e826d7ddf3c85920")

# ------------------------------------------------------------------ PITBOSS log poisoning (CVE-2026-88771)
# Stage 1: a base64 payload is parked in the User-Agent of a web request (httpaccess*.log). Stage 2: a failed login
# whose user name imitates a pitboss crash record ("pitboss PPE missed too many heartbeatsNSPPE;<command>",
# "pitboss PPE unexpectedly died NSPPE-00;<command>;# X") is written to ns.log. Stage 3: the Call Home crash checker
# (/netscaler/ns_monuploadd_err.pl) reads ns.log, takes the text after NSPPE and runs it with sh or php, usually
# pulling the stage 1 payload back out of the access log. Genuine pitboss records come from the pitboss process itself
# ("pitboss[26]: ... NSPPE-01 (11854) unexpectedly died") and never carry a command, so they are hunted separately (C, D).
_SHELL_META = r"(?:[;`|]|\$\(|\$\{|&&|%3b|%60|%7c|%24%28|%24%7b|%26%26)"
PITBOSS_FORGED_RX = re.compile(
    r"NSPPE(?:-\d+)?\s{0,2}" + _SHELL_META                                   # a command right after NSPPE / NSPPE-00
    + r"|pitboss.{0,200}?PPE.{0,80}?(?:`|\$\(|\$\{)"                         # pitboss text with `...`, $(...) or ${IFS}
    + r"|pitboss.{0,200}?(?:missed too many heartbeats|unexpectedly died).{0,200}?" + _SHELL_META
    + r"|(?:\buser(?:name)?\s*:?\s*<?|\bcontext\s+|\bfor\s+<|packet engine for\s+)(?:\S{1,12}\s+)?pitboss",  # in a user name
    re.I)
# Lines that quote the pattern instead of carrying it: CLI command audit records in ns.log (a responder's
# shell "grep ..." run from the NetScaler CLI) and regular expressions written by people hunting for it.
PITBOSS_QUOTE_RX = re.compile(r"CMD_EXECUTED|pitboss\.\*|NSPPE\.\*", re.I)
CALLHOME_INJECT_RX = re.compile(
    r"NSPPE(?:-\d+)?\s{0,2}" + _SHELL_META + r"|NSPPE\S{0,8}\s{0,2}[^\s`$]{0,40}(?:`|\$\(|\$\{)", re.I)
# Text of the injected command: after NSPPE; up to the attacker's comment (;# X) or the end of the logged user name.
INJECT_CMD_RX = re.compile(
    r"NSPPE(?:-\d+)?\s{0,2};(?P<cmd>.*?)(?=\s*;\s*#[\s+-]*X\b|\s*#[\s+]*-[\s+]*X\b|>\s*,\s*factor\b|>\"?\s*$|\"\s*$"
    r"|\s+-\s+Client_ip\b|\s*,\s*trying to find\b|$)", re.I)
INJECT_SUBST_RX = re.compile(r"`([^`]{1,800})`|\$\((.{1,800}?)\)")
CLIENT_IP_RX = re.compile(r"Client_?ip[:\s]+([0-9A-Fa-f.:]{3,})|@(\d{1,3}(?:\.\d{1,3}){3})\b", re.I)
# Stage 1 payloads: base64 after INDEX: or K: (as reported by Zscaler), base64 of "<?php" (PD9waHA), or any long base64
# run with upper case, lower case and digits. Only the User-Agent field is examined.
UA_PAYLOAD_RX = re.compile(r"(?P<mark>INDEX:|\bK:)(?P<pre>[A-Za-z0-9+/=]{16,})|(?P<run>[A-Za-z0-9+/]{60,}={0,2})"
                           r"|(?P<php>PD9waHA[A-Za-z0-9+/]*={0,2})")
# SHA-256 published by Unit 42 (2026-09-30): the stage 1 base64 text, the shell script it decodes to, and the nsg64.deb
# web shell. Staged User-Agent payloads (B2) and files in the web, client download and temp folders are compared.
KNOWN_SHA256 = {
    "1bd314b661396c7086f6367fbbb48025e03ca2de69c073d53a8b0a38aa5fbb7d": "Unit 42 published hash (base64 payload)",
    "79c65fa04541032e251fa4796b97800374b63c7982593dd1a2e0db605d429186": "Unit 42 published hash (decoded shell script)",
    "ae22ef2517b5c0fb47f78745b9cb5260acee0e751b89bcd354640ff8bc8d29ec": "Unit 42 published hash (nsg64.deb web shell)",
}


class PitbossMatcher:
    """Hunt B: forged pitboss / NSPPE records. A substring check skips the regex on almost every line."""

    def search(self, text):
        low = text.lower()
        if "nsppe" not in low and "pitboss" not in low:
            return None
        if PITBOSS_QUOTE_RX.search(text):
            return None
        return PITBOSS_FORGED_RX.search(text)


class CallHomeMatcher:
    """Hunt A: Call Home crash checker lines that carry an injected command."""

    def search(self, text):
        if "nsppe" not in text.lower():
            return None
        return CALLHOME_INJECT_RX.search(text)


def user_agent(text):
    """The last quoted field of an access-log line (the User-Agent); '' when the line has fewer than two."""
    s = text.rstrip()
    if s.count('"') < 4:
        return ""
    j = s.rfind('"')
    i = s.rfind('"', 0, j)
    return s[i + 1:j] if i >= 0 else ""


B64_TOKEN_RX = re.compile(r"[A-Za-z0-9+/]{8,}={0,2}")
CMDISH_RX = re.compile(r"://|<\?php|/bin/|/tmp/|\b(?:curl|wget|fetch|bash|php|perl|python[23]?|echo|chmod|base64|whoami|"
                       r"uname|mkdir|nohup|busybox|tftp|cat|tar)\b|(?:^|[|;&]\s*)(?:sh|id|ls)\b")


def readable_b64(tok):
    """True when a whole token is base64 (8+ characters) that decodes to readable ASCII text with a command, URL or
    PHP code in it. Ordinary user-agent words almost never decode to printable ASCII, let alone to a command."""
    if not B64_TOKEN_RX.fullmatch(tok):
        return False
    raw = decode_b64(tok)
    if len(raw) < 6:
        return False
    try:
        txt = raw.decode("ascii")
    except UnicodeDecodeError:
        return False
    good = sum(1 for ch in txt if " " <= ch <= "~" or ch in "\r\n\t")
    return good / len(txt) >= 0.95 and bool(CMDISH_RX.search(txt))


def ua_payload(ua):
    """(marker, base64 text) of a payload staged in a User-Agent, or None."""
    for m in UA_PAYLOAD_RX.finditer(ua):
        if m.group("pre"):
            return m.group("mark"), m.group("pre")
        if m.group("php"):
            return "PD9waHA (base64 of <?php)", m.group("php")
        run = m.group("run")
        if re.search(r"[A-Z]", run) and re.search(r"[a-z]", run) and re.search(r"\d", run):
            return "long base64 run", run
    for tok in ua.split():     # short payloads: a whole word of the User-Agent that decodes to a readable command
        if readable_b64(tok):
            return "base64 word that decodes to text", tok
    return None


class UaPayloadMatcher:
    """Hunt B2: web requests whose User-Agent carries a base64 payload (PITBOSS stage 1)."""

    def search(self, text):
        ua = user_agent(text)
        if len(ua) < 8:
            return None
        # quick exit for ordinary agents: no marker, no stretch long enough for a 60-character run, and no word made
        # only of base64 characters
        words = ua.split()
        if "INDEX:" not in ua and "K:" not in ua and "PD9waHA" not in ua and not any(len(w) >= 60 for w in words) \
                and not any(len(w) >= 8 and "." not in w and B64_TOKEN_RX.fullmatch(w) for w in words):
            return None
        return ua_payload(ua)


def decode_b64(s):
    """Bytes of a base64 text, tolerating missing padding; b'' when it is not base64."""
    s = re.sub(r"[^A-Za-z0-9+/]", "", s)
    if len(s) % 4 == 1:
        s = s[:-1]
    try:
        return base64.b64decode(s + "=" * (-len(s) % 4))
    except (binascii.Error, ValueError):
        return b""


def payload_kind(raw):
    """What a decoded payload is, without keeping its content."""
    if not raw:
        return "not valid base64"
    txt = raw[:8192].decode("utf-8", "replace")
    if "<?php" in txt or "<?=" in txt:
        return "PHP code"
    good = sum(1 for ch in txt if ch.isprintable() or ch in "\r\n\t")
    if good / len(txt) < 0.85:
        return f"binary ({len(raw)} bytes)"
    if txt.startswith("#!") or re.search(r"\b(curl|wget|fetch|sh|bash|echo|chmod|mkdir|base64|perl|python3?|php|tar|nc|"
                                         r"uname|whoami|id|cat|ls|ps)\b", txt):
        return "shell commands"
    return "text"


# Hunt sections: file name filters ("!" excludes) and patterns (case-insensitive) or matcher objects. Only files
# directly in /var/log are read, except for H, which reads every exported text log. H is built from BUILTIN_IOC_IPS
# plus --ioc-ip. Shell history and audit logs are left out of B: responders' own grep commands quote the pattern.
SECTIONS = [
    ("A", "A_CallHome_Exec", "PITBOSS stage 3: Call Home crash checker lines that carry an injected command",
     ["callhomedebug*"], CallHomeMatcher()),
    ("B", "B_Pitboss_Log_Poisoning", "PITBOSS stage 2: forged pitboss / NSPPE records (CVE-2026-88771 log poisoning)",
     ["*", "!callhomedebug*", "!notice.log*", "!sh.log*", "!bash.log*"], PitbossMatcher()),
    ("B2", "B2_Staged_UA_Payloads", "PITBOSS stage 1: base64 payloads staged in the User-Agent", ["httpaccess*"],
     UaPayloadMatcher()),
    ("C", "C_nsaaad_Crash_Reboot", "nsaaad crashes, restart limits and pitboss reboots", ["messages*", "ns.log*"],
     r"nsaaad \(\d+\) unexpectedly died|Pitboss rebooted the system|proc nsaaad.*(?:SIGNALED|EXITED)"
     r"|maximum number of restarts|Pitboss declaring system failure|All monitored processes have exited, rebooting"),
    ("D", "D_NSPPE_Crashes", "Genuine packet engine (NSPPE) crash records from pitboss", ["messages*", "ns.log*"],
     r"pitboss.*NSPPE-\d+ \(\d+\) (?:unexpectedly died|missed too many heartbeats)|pitboss.*NOT restarting NSPPE"),
    ("E", "E_Webshell_Requests", "Requests to the webshell paths", ["httpaccess*"], r"/logon/LogonPoint/custom/(receiver|\.ctxs|\.slap)"),
    ("E2", "E2_Other_Probe_Paths", "Requests to other implant or probe paths", ["httpaccess*"],
     r"receiver\.min|/vpn/media/[^ ]+\.ico|:9910|/logon/LogonPoint/Authentication/GetUserName"
     r"|/vpn/scripts/linux/(?:nsgclient18|nsgser18|nsgsupport|nsgpackage64|nsgbuild|nsg64)\.deb\b"),
    ("F", "F_Relay_Installer", ".relay installer lines in the shell audit log", ["notice.log*"],
     r"/nsconfig/\.relay|relay\.env|target_agent|ssh -N -R"),
    ("G", "G_Implant_Cron_Runs", "Cron runs of the implant", ["cron*"], r"/nsconfig/\.(slap|relay)/"),
    ("H", "H_Attacker_IP_Lines", "Log lines that mention the attacker IPs", ["*"], None),
    ("I", "I_Shell_Suspicious", "Suspicious commands in the shell history logs", ["sh.log*", "bash.log*"],
     r"relay|receiver|loot|curl|fetch |crontab|truncat|: >|cat /dev/null|rm -|nsshutdown|reboot|\.slap|\.ux|python3|ssh |scp "
     r"|tar |webhook"),
]
SECTION_I_EXCLUDE = re.compile(r"boot\.sh|/var/tmp/\.ux/agent\.log|CRON_TMP", re.I)

# Discovery: hidden files that a stock NetScaler has, and folders that change normally.
BENIGN_HIDDEN_RX = re.compile(
    r"^/flash/(\.snap|\.version|\.nsva|\.testing)$"
    r"|^/flash/nsconfig/\.(eula_run|nshostuuid|callhome\.conf|skf|rpcs|F1\.key|F2\.key|nscfg_hostuid|telemetry\.conf)$"
    r"|^/var/(\.snap|\.monit\.id|\.monit\.state|\.agent)$|/\.signedexe\.manifest$|^/var/nsinstall/\.ns\.version$"
    r"|^/var/run/\.(NETSCALER|NSCONFIG|NSVERSION|NSKERNEL)$|/\.ssh(/|$)|/\.nscli_history$|/\.nsclirc$"
    r"|^/var/tmp/\.shrun\.cache$|^/var/netscaler/gui/admin_ui/php/system/\.htaccess$|^/var/nsinstall/build_[\d.]+/\.ns\.version$"
    r"|^/var/netscaler/help/.*/\.DS_Store$"
    r"|^/var/(mastools|perl5|python|nextgen|pgsql)/|^/var/log/\.cache/|^/flash/nsconfig/license/")
FLASH_ROTATING_RX = re.compile(
    r"^/flash/nsconfig/(ns\.conf(\..*)?|ZebOS\.conf(\..*)?|admparam\.conf|unified\.conf|snmpd\.conf|pooledlicense\.conf"
    r"|admautoreg\.state|cfgsync_known_hosts|fsync_known_hosts|\.nscfg_hostuid|\.telemetry\.conf|updated_.*"
    r"|license/.*|monitors/.*|loginschema/.*|apispec/.*|ssl/.*)$|^/flash/(ns-[\d.-]+\.gz|secboot/.*|boot/.*|\.version)$")
TMP_NOISE_RX = re.compile(
    r"^/var/(tmp|nstmp)/(par-[0-9a-f]+(/|$)|log-\d{4}-\d\d-\d\d\.php$|[\w.]*_(bios|bmc)\.bin$|(bios|bmc)_releases$"
    r"|Fortville_Silicom_Intel(/|$)|Mellanox(/|$)|sum$|install$|tmp[\w]{8}$|ch_metrics\.txt$)")
WEB_NOISE_RX = re.compile(r"^/var/(run|netscaler/gui|netscaler/nitro|vpn/themes|app_catalog|nsinstall)/")
# VPN and EPA client download folders (served as /vpn/scripts/linux/ and /epa/scripts/linux/). Unit 42 (2026-09-30)
# reported the actor hosting webshells there under package-like names. A real .deb starts with "!<arch>\n" and a real
# .rpm with ED AB EE DB; anything else with those extensions is a disguised file.
CLIENT_PKG_DIR_RX = re.compile(r"^/var/netscaler/gui/(vpn|epa)/scripts/linux/[^/]+$")
UNIT42_PKG_NAMES = {"nsgclient18.deb", "nsgser18.deb", "nsgsupport.deb", "nsgpackage64.deb", "nsgbuild.deb", "nsg64.deb"}
PKG_MAGIC = {".deb": b"!<arch>\n", ".rpm": b"\xed\xab\xee\xdb"}
# Web-served folders on /var. Fox-IT only inspects *.php files there; this actor hid PHP web shells in other files
# (.ctxs.receiver, package-named .deb files) and reached them through aliases. In the GUI tree, which holds the stock
# PHP application, only static file types, extension-less files and hidden files are read; elsewhere every non-PHP file.
WEB_ROOT_RX = re.compile(r"^/var/(netscaler/logon|netscaler/gui|netscaler/ns_gui|vpn)/")
GUI_ROOT_RX = re.compile(r"^/var/netscaler/(gui|ns_gui)/")
PHP_EXTS = {".php", ".php3", ".php4", ".php5", ".php7", ".phtml", ".phps", ".inc", ".tpl"}
STATIC_EXTS = {".css", ".js", ".map", ".json", ".xml", ".txt", ".png", ".gif", ".jpg", ".jpeg", ".ico", ".svg", ".woff",
               ".woff2", ".ttf", ".eot", ".deb", ".rpm", ".gz", ".zip", ".receiver", ".min"}
PHP_CODE_RX = re.compile(rb"<\?php|<\?=\s*\$", re.I)
# Files hashed and compared with KNOWN_SHA256: web folders outside the GUI tree, the client download folders, the
# temp folders, and any file another check flagged. Larger files are skipped (the published files are small).
HASH_SCOPE_RX = re.compile(r"^/var/(netscaler/logon|vpn|tmp|nstmp)/|^/var/netscaler/(gui|ns_gui)/(vpn|epa)/scripts/linux/")
HASH_MAX_BYTES = 20 << 20
LOGON_DROP_DIR = "/var/netscaler/logon/LogonPoint/custom"
WALK_SKIP_WINDOW = ("/var/log/", "/var/nslog/", "/var/nstrace/", "/var/core/", "/var/crash/", "/var/netscaler/locdb/",
                    "/var/mastools/logs/")

UTX_TYPES = {0: "EMPTY", 1: "BOOT_TIME", 2: "OLD_TIME", 3: "NEW_TIME", 4: "USER_PROCESS", 5: "INIT_PROCESS",
             6: "LOGIN_PROCESS", 7: "DEAD_PROCESS", 8: "SHUTDOWN_TIME"}
UTX_LOG_RX = re.compile(r"^utx\.log(\.\d+)?(\.gz)?$")
UTX_FIXED_RX = re.compile(r"^utx\.(lastlogin|active)$")

MONTHS = {m: i for i, m in enumerate("Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split(), 1)}
SYSLOG_TS = re.compile(r"^([A-Z][a-z]{2}) +(\d{1,2}) (\d\d):(\d\d):(\d\d)")
NSGMT_TS = re.compile(r"(\d\d)/(\d\d)/(\d{4}):(\d\d):(\d\d):(\d\d) GMT")
HTTP_TS = re.compile(r"\[(\d\d)/([A-Z][a-z]{2})/(\d{4}):(\d\d):(\d\d):(\d\d) ([+-]\d{4})\]")
HTTPACC_RX = re.compile(
    r'^(?P<ip>\S+)(?: -> (?P<via>\S+))? \S+ \S+ \[(?P<ts>[^\]]+)\] \[(?P<pid>\d+)\] "(?P<req>[^"]*)" (?P<status>\d{3}|-) '
    r'(?P<bytes>\S+)(?: "(?P<ref>[^"]*)")?(?: "(?P<ua>[^"]*)")?')
IPV4_RX = re.compile(r"(?<![\d.])(\d{1,3}(?:\.\d{1,3}){3})(?![\d.])")
URL_RX = re.compile(r"(?i)\b((?:https?|ftp)://[^\s'\"<>;|]+)")

LOG = None


# ------------------------------------------------------------------ small helpers
def say(msg):
    line = f"[{dt.datetime.now().strftime('%H:%M:%S')}] {msg}"
    print(line, flush=True)
    if LOG:
        LOG.write(line + "\n")
        LOG.flush()


def now_utc():
    return dt.datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")


def ts(x):
    """epoch seconds or datetime -> 'YYYY-MM-DD HH:MM:SS' UTC ('' when unknown)."""
    if x is None or x == "" or x == 0:
        return ""
    try:
        if isinstance(x, dt.datetime):
            if x.tzinfo is None:
                x = x.replace(tzinfo=UTC)
            return x.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S")
        return dt.datetime.fromtimestamp(float(x), UTC).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ""


def parse_ts(s):
    try:
        return dt.datetime.strptime(s, "%Y-%m-%d %H:%M:%S").replace(tzinfo=UTC)
    except Exception:
        return None


def lp(p):
    """Path string usable by Windows APIs even when longer than 260 characters."""
    s = os.path.abspath(str(p))
    if os.name == "nt" and not s.startswith("\\\\?\\") and len(s) > 240:
        s = "\\\\?\\" + s
    return s


_BAD = re.compile(r'[<>:"|?*\\\x00-\x1f]')
_RESERVED = re.compile(r"^(con|prn|aux|nul|com[1-9]|lpt[1-9])(\..*)?$", re.I)


def safe_part(name):
    n = _BAD.sub("_", name).rstrip(" .") or "_"
    if _RESERVED.match(n):
        n = "_" + n
    return n


def safe_rel(image_rel):
    return Path(*[safe_part(x) for x in str(image_rel).strip("/").split("/") if x])


def sha256_file(p):
    h = hashlib.sha256()
    with open(lp(p), "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


_NULRUN = re.compile("\x00+")
_CTRL = re.compile("[\x01-\x08\x0b\x0c\x0e-\x1f\x7f\ufffe\uffff]")


def xl(s):
    """Text that Excel accepts: NUL runs shown as [NUL xN], other control characters as \\xNN, long cells cut."""
    if s is None:
        return ""
    s = str(s)
    if "\x00" in s:
        s = _NULRUN.sub(lambda m: f"[NUL x{len(m.group(0))}]", s)
    if _CTRL.search(s):
        s = _CTRL.sub(lambda m: ("\\x%02x" if ord(m.group(0)) < 256 else "\\u%04x") % ord(m.group(0)), s)
    if len(s) > XL_MAX_CELL:
        s = s[:XL_MAX_CELL] + f" ...[cut: {len(s)} characters in total, full text in the exported log]"
    return s


def line_time(text, ref):
    """Best-effort UTC time of a log line ('' if none). ref = datetime giving the year for syslog lines."""
    m = NSGMT_TS.search(text)
    if m:
        mo, d, y, H, M, S = m.groups()
        return f"{y}-{mo}-{d} {H}:{M}:{S}"
    m = HTTP_TS.search(text)
    if m:
        d, mon, y, H, M, S, off = m.groups()
        try:
            t = dt.datetime(int(y), MONTHS.get(mon, 1), int(d), int(H), int(M), int(S))
            sign = 1 if off[0] == "+" else -1
            t -= sign * dt.timedelta(hours=int(off[1:3]), minutes=int(off[3:5]))
            return t.strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return ""
    m = SYSLOG_TS.match(text)
    if m and ref is not None:
        mon, d, H, M, S = m.groups()
        mo = MONTHS.get(mon)
        if mo:
            y = ref.year if (mo, int(d)) <= (ref.month, ref.day) else ref.year - 1
            return f"{y:04d}-{mo:02d}-{int(d):02d} {H}:{M}:{S}"
    return ""


def iter_lines(path):
    """Yields raw byte lines of a plain or .gz file. A truncated .gz ends quietly; see the inventory note."""
    opener = gzip.open if str(path).lower().endswith(".gz") else open
    with opener(lp(path), "rb") as fh:
        for raw in fh:
            yield raw


def read_head(path, n=65536):
    opener = gzip.open if str(path).lower().endswith(".gz") else open
    with opener(lp(path), "rb") as fh:
        return fh.read(n)


def looks_binary(data):
    if not data:
        return False
    sample = data[:65536]
    printable = sum(1 for b in sample if 32 <= b < 127 or b in (9, 10, 13) or b >= 128)
    return printable / len(sample) < 0.85


def pkg_version(name):
    try:
        from importlib.metadata import version
        return version(name)
    except Exception:
        return ""


# ------------------------------------------------------------------ workbook writer
class Book:
    """Streams rows into an .xlsx (XlsxWriter constant-memory). Strings are never turned into formulas or links."""

    def __init__(self, path):
        import xlsxwriter
        self.path = path
        self.wb = xlsxwriter.Workbook(lp(path), {"constant_memory": True, "strings_to_numbers": False,
                                                 "strings_to_formulas": False, "strings_to_urls": False})
        self.hdr = self.wb.add_format({"bold": True, "bg_color": "#DDEBF7", "border": 1})
        self.bold = self.wb.add_format({"bold": True})
        self.wrap = self.wb.add_format({"text_wrap": True, "valign": "top"})
        self.names = set()

    def _name(self, name):
        n = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
        base, i = n, 2
        while n.lower() in self.names:
            suffix = f" {i}"
            n = base[:31 - len(suffix)] + suffix
            i += 1
        self.names.add(n.lower())
        return n

    def table(self, name, headers, rows, widths=None):
        state = {"ws": None, "r": 0, "part": 0}
        ncol = len(headers)

        def open_sheet():
            state["part"] += 1
            ws = self.wb.add_worksheet(self._name(name if state["part"] == 1 else f"{name} {state['part']}"))
            for c, h in enumerate(headers):
                ws.write_string(0, c, h, self.hdr)
                ws.set_column(c, c, widths[c] if widths and c < len(widths) else 18)
            ws.freeze_panes(1, 0)
            state["ws"], state["r"] = ws, 0

        open_sheet()
        n = 0
        for row in rows:
            if state["r"] >= XL_MAX_ROWS:
                state["ws"].autofilter(0, 0, state["r"], ncol - 1)
                open_sheet()
            state["r"] += 1
            ws, r = state["ws"], state["r"]
            for c, v in enumerate(row):
                if v is None or v == "":
                    continue
                if isinstance(v, bool):
                    ws.write_string(r, c, "yes" if v else "no")
                elif isinstance(v, (int, float)):
                    ws.write_number(r, c, v)
                else:
                    ws.write_string(r, c, xl(v))
            n += 1
        state["ws"].autofilter(0, 0, max(state["r"], 1), ncol - 1)
        return n

    def notes(self, name, lines, width=140):
        ws = self.wb.add_worksheet(self._name(name))
        ws.set_column(0, 0, width)
        for r, line in enumerate(lines):
            if isinstance(line, tuple):
                ws.write_string(r, 0, xl(line[0]), self.bold)
            else:
                ws.write_string(r, 0, xl(line), self.wrap)

    def close(self):
        self.wb.close()


# ------------------------------------------------------------------ disk discovery
def _parse_descriptor_text(text):
    d = {"cid": "", "parent_cid": "", "create_type": "", "parent_hint": "", "extents": []}
    for line in text.splitlines():
        s = line.strip()
        m = re.match(r'^(CID|parentCID|createType|parentFileNameHint)\s*=\s*"?([^"]*)"?', s)
        if m:
            key = {"CID": "cid", "parentCID": "parent_cid", "createType": "create_type",
                   "parentFileNameHint": "parent_hint"}[m.group(1)]
            d[key] = m.group(2).strip()
            continue
        m = re.match(r'^(RW|RDONLY|NOACCESS)\s+(\d+)\s+(\w+)\s*(?:"([^"]*)")?', s)
        if m:
            d["extents"].append({"access": m.group(1), "sectors": int(m.group(2)), "type": m.group(3),
                                 "file": m.group(4) or ""})
    return d


def read_descriptor(path):
    """Returns a descriptor dict for a VMDK descriptor (text or embedded in a KDMV sparse file), else None."""
    try:
        size = path.stat().st_size
        with open(lp(path), "rb") as fh:
            head = fh.read(4096)
            if head[:4] == b"KDMV":
                desc_off, desc_size = struct.unpack("<QQ", head[28:44])
                if not desc_off or not desc_size:
                    return {"cid": "", "parent_cid": "", "create_type": "monolithicSparse", "parent_hint": "",
                            "extents": [], "embedded": True}
                fh.seek(desc_off * 512)
                text = fh.read(desc_size * 512).split(b"\x00")[0].decode("latin-1", "replace")
                d = _parse_descriptor_text(text)
                d["embedded"] = True
                return d
            if size > 65536:
                return None
            fh.seek(0)
            text = fh.read().decode("latin-1", "replace")
        if "createType" not in text:
            return None
        d = _parse_descriptor_text(text)
        d["embedded"] = False
        return d
    except Exception:
        return None


def parse_vmx(path):
    info = {"vmx": str(path), "display_name": "", "guest_os": "", "disks": []}
    try:
        for line in open(lp(path), encoding="utf-8", errors="replace"):
            m = re.match(r'^\s*([\w:.]+)\s*=\s*"(.*)"\s*$', line)
            if not m:
                continue
            k, v = m.group(1), m.group(2)
            if k == "displayName":
                info["display_name"] = v
            elif k == "guestOS":
                info["guest_os"] = v
            elif re.match(r"^(scsi|sata|nvme|ide)\d+:\d+\.fileName$", k) and v.lower().endswith(".vmdk"):
                info["disks"].append((k.split(".")[0], v))
    except Exception as ex:
        info["error"] = repr(ex)
    return info


def parse_vmsd(path):
    snaps, cur = {}, ""
    try:
        for line in open(lp(path), encoding="utf-8", errors="replace"):
            m = re.match(r'^\s*([\w.]+)\s*=\s*"(.*)"\s*$', line)
            if not m:
                continue
            k, v = m.groups()
            m2 = re.match(r"^(snapshot\d+)\.(displayName|createTimeHigh|createTimeLow|uid|filename)$", k)
            if m2:
                snaps.setdefault(m2.group(1), {})[m2.group(2)] = v
            elif k == "snapshot.current":
                cur = v
    except Exception:
        pass
    out = []
    for s in snaps.values():
        try:
            us = (int(s.get("createTimeHigh", "0")) << 32) + (int(s.get("createTimeLow", "0")) & 0xFFFFFFFF)
            created = ts(us / 1e6) if us else ""
        except ValueError:
            created = ""
        out.append({"name": s.get("displayName", ""), "created_utc": created, "uid": s.get("uid", ""),
                    "current": s.get("uid", "") == cur})
    return out


RAW_IMAGE_EXTS = (".dd", ".raw", ".img")
ARCHIVE_EXTS = (".zip", ".7z", ".rar", ".ova", ".tar", ".tgz", ".dd.gz", ".raw.gz", ".img.gz", ".vmdk.gz", ".dd.xz",
                ".raw.xz", ".img.xz", ".dd.bz2", ".raw.bz2", ".img.bz2")
MIN_IMAGE_BYTES = 1 << 20          # when a folder is searched, smaller files are not treated as disk images
_EWF_SEGMENT = re.compile(r"^\.[E-Ze-z][0-9A-Za-z]{2}$")   # E01..E99, EAA..ZZZ (the names dissect looks for)
_EWF_NUMBERED = re.compile(r"^\.[Ee]\d{2}$")                 # E01..E99
OTHER_IMAGE_FORMATS = {".vhd": "VHD", ".vhdx": "VHDX", ".qcow2": "QCOW2", ".vdi": "VDI"}


def _file_size(p):
    try:
        return p.stat().st_size
    except OSError:
        return 0


def split_parts(first):
    """Parts of a split raw image (same name, numeric extension), in the order dissect reads them, plus problems."""
    parts = sorted((x for x in first.parent.iterdir()
                    if x.is_file() and x.stem == first.stem and x.suffix[1:].isdigit()), key=lambda x: x.name)
    nums = [int(x.suffix[1:]) for x in parts]
    problems = []
    if nums != sorted(nums):
        problems.append("the split parts are not zero-padded (for example .1, .2, .10), so they would be read out of "
                        "order; rename them to .001, .002, ...")
    if nums:
        missing = sorted(set(range(min(nums), max(nums) + 1)) - set(nums))
        if missing:
            width = len(first.suffix) - 1
            problems.append("split parts missing: " + ", ".join(f".{n:0{width}d}" for n in missing))
    return parts, problems


def ewf_parts(first):
    """Segments of an E01 set (E01, E02, ... EAA ...), in the order dissect reads them, plus problems."""
    parts = sorted((x for x in first.parent.iterdir()
                    if x.is_file() and x.stem == first.stem and _EWF_SEGMENT.match(x.suffix)), key=lambda x: x.name)
    problems = []
    nums = sorted(int(x.suffix[2:]) for x in parts if x.suffix[2:].isdigit())
    if nums:
        missing = sorted(set(range(1, max(nums) + 1)) - set(nums))
        if missing:
            problems.append("E01 segments missing: " + ", ".join(f".E{n:02d}" for n in missing))
    for x in parts:
        try:
            with open(lp(x), "rb") as fh:
                magic = fh.read(3)
        except OSError as ex:
            problems.append(f"{x.name}: {ex}")
            continue
        if magic != b"EVF":
            problems.append(f"{x.name} has the same name as the E01 set but is not an E01 segment; move it to another "
                            "folder, or the E01 reader will fail")
    return parts, problems


def image_cand(first, fmt, parts, problems):
    """A disk candidate for a non-VMDK image (raw, split raw or E01)."""
    segs = [{"file": x.name, "size": _file_size(x)} for x in parts]
    total = sum(s["size"] for s in segs)
    problems = list(problems)
    if fmt in ("raw", "split raw") and total % 512:
        problems.append(f"{first.name}: {total} bytes is not a whole number of 512-byte sectors (truncated image?)")
    if len(parts) > 1:
        fmt = f"{fmt} ({len(parts)} {'segments' if fmt == 'E01' else 'parts'})"
    return {"label": "", "folder": str(first.parent), "leaf": str(first), "base": str(first), "chain": [str(first)],
            "chain_desc": [], "extents": [], "problems": problems, "vmx": [], "vmsd": [], "format": fmt,
            "segments": segs, "image_bytes": total}


def _assign_labels(cands, single_file=False):
    """Label = folder name, or folder_file when a folder holds more than one image; made unique (case-insensitive)."""
    per_folder = collections.Counter(c["folder"] for c in cands)
    used = set()
    for c in cands:
        folder, stem = Path(c["folder"]).name, Path(c["leaf"]).stem
        label = safe_part(folder if per_folder[c["folder"]] == 1 and not single_file else f"{folder}_{stem}")
        unique, n = label, 1
        while unique.lower() in used:
            n += 1
            unique = f"{label}_{n}"
        used.add(unique.lower())
        c["label"] = unique
    return cands


def discover(root, max_depth=3):
    """Finds disk images under root (a folder) or opens root itself (a file). Returns (candidates, notes).

    Folders are searched max_depth levels deep for VMDK chains, raw images (.dd .raw .img), split raw images
    (.000/.001 with numbered parts) and E01 sets. A raw file that a VMDK descriptor uses as its data file is
    analyzed through the descriptor only."""
    root = Path(root)
    if root.is_file():
        return discover_file(root)
    notes = []
    descs, vmx_files, vmsd_files, raw_files, numbered, ewf_firsts = {}, [], [], [], {}, []
    for dirpath, dirnames, filenames in os.walk(root):
        depth = len(Path(dirpath).relative_to(root).parts)
        if depth >= max_depth:
            dirnames[:] = []
        dirnames.sort()
        for fn in sorted(filenames):
            p = Path(dirpath) / fn
            low = fn.lower()
            if low.endswith(".vmdk"):
                d = read_descriptor(p)
                if d is not None:
                    descs[p] = d
            elif low.endswith(".vmx"):
                vmx_files.append(p)
            elif low.endswith(".vmsd"):
                vmsd_files.append(p)
            elif low.endswith(RAW_IMAGE_EXTS):
                raw_files.append(p)
            elif low.endswith(".e01"):
                ewf_firsts.append(p)
            elif low.endswith(".ex01"):
                notes.append(f"Not opened (EWF2 .Ex01 is not supported; convert it to E01 or raw first): {p}")
            elif p.suffix[1:].isdigit() and len(p.suffix) >= 4:
                numbered.setdefault((p.parent, p.stem), []).append(p)
            elif low.endswith(ARCHIVE_EXTS) and depth <= 1:
                notes.append(f"Not opened (archive; extract it first): {p}")
    by_folder = collections.defaultdict(dict)
    for p, d in descs.items():
        by_folder[p.parent][p.name.lower()] = (p, d)
    cands = []
    used_as_extent = set()
    for folder, items in sorted(by_folder.items(), key=lambda x: str(x[0])):
        extent_names = {e["file"].lower() for _, d in items.values() for e in d["extents"] if e["file"]}
        used_as_extent |= {(folder, n) for n in extent_names}
        heads = {k: v for k, v in items.items() if k not in extent_names}
        parents = {Path(d["parent_hint"].replace("\\", "/")).name.lower() for _, d in heads.values() if d["parent_hint"]}
        leaves = [v for k, v in heads.items() if k not in parents]
        vmx = [parse_vmx(x) for x in vmx_files if x.parent == folder]
        vmsd = [s for x in vmsd_files if x.parent == folder for s in parse_vmsd(x)]
        for leaf_path, leaf in sorted(leaves, key=lambda x: x[0].name):
            chain, cur, seen = [], (leaf_path, leaf), set()
            while cur and cur[0].name.lower() not in seen:
                seen.add(cur[0].name.lower())
                chain.append(cur)
                hint = cur[1]["parent_hint"]
                if not hint or cur[1]["parent_cid"].lower() == "ffffffff":
                    break
                nxt = heads.get(Path(hint.replace("\\", "/")).name.lower())
                if nxt is None:
                    notes.append(f"{cur[0].name}: parent disk '{hint}' is not in {folder}")
                    break
                cur = nxt
            cands.append(vmdk_cand(folder, chain, vmx, vmsd))
    for p in raw_files:
        if (p.parent, p.name.lower()) in used_as_extent:
            notes.append(f"{p} is the data file of a VMDK descriptor in the same folder; it is analyzed through that "
                         "descriptor")
        elif _file_size(p) < MIN_IMAGE_BYTES:
            notes.append(f"Skipped (smaller than 1 MiB, so not a disk image): {p}")
        else:
            cands.append(image_cand(p, "raw", [p], []))
    for (folder, stem), files in sorted(numbered.items(), key=lambda x: str(x[0])):
        nums = sorted(int(x.suffix[1:]) for x in files)
        if nums[0] > 1:
            continue                                   # no .000/.001 first part: not the start of a split image
        parts, problems = split_parts(files[0])
        first = parts[0]
        if any((folder, x.name.lower()) in used_as_extent for x in parts):
            notes.append(f"{first} is a data file of a VMDK descriptor in the same folder; it is analyzed through that "
                         "descriptor")
            continue
        if sum(_file_size(x) for x in parts) < MIN_IMAGE_BYTES:
            notes.append(f"Skipped (smaller than 1 MiB, so not a disk image): {first}")
            continue
        if problems:
            notes.append(f"Not opened ({'; '.join(problems)}): {first}")
            continue
        cands.append(image_cand(first, "split raw" if len(parts) > 1 else "raw", parts, []))
    for p in ewf_firsts:
        parts, problems = ewf_parts(p)
        if problems:
            notes.append(f"Not opened ({'; '.join(problems)}): {p}")
            continue
        cands.append(image_cand(p, "E01", parts, []))
    cands.sort(key=lambda c: (c["folder"].lower(), Path(c["leaf"]).name.lower()))
    return _assign_labels(cands), notes


def vmdk_cand(folder, chain, vmx, vmsd):
    """A disk candidate for a VMDK chain (leaf first, base last)."""
    ext_rows, problems = [], []
    for dp, dd in chain:
        for e in dd["extents"]:
            ef = dp.parent / e["file"] if e["file"] else None
            exists = bool(ef and ef.exists())
            size = ef.stat().st_size if exists else 0
            expected = e["sectors"] * 512 if e["type"].upper() in ("VMFS", "FLAT") else None
            status = "ok"
            if not exists:
                status = "MISSING"
            elif expected and size < expected:
                status = f"INCOMPLETE ({size / 2**30:.1f} of {expected / 2**30:.1f} GiB)"
            if status != "ok":
                problems.append(f"{e['file']}: {status}")
            ext_rows.append({"descriptor": dp.name, "extent": e["file"], "type": e["type"],
                             "sectors": e["sectors"], "size": size, "status": status})
    fmt = "VMDK" if len(chain) == 1 else f"VMDK (snapshot chain of {len(chain)} disks)"
    return {"label": "", "folder": str(folder), "leaf": str(chain[0][0]),
            "base": str(chain[-1][0]), "chain": [str(x[0]) for x in chain],
            "chain_desc": [{"file": x[0].name, "cid": x[1]["cid"], "parent_cid": x[1]["parent_cid"],
                            "type": x[1]["create_type"], "parent_hint": x[1]["parent_hint"]} for x in chain],
            "extents": ext_rows, "problems": problems, "vmx": vmx, "vmsd": vmsd, "format": fmt, "segments": []}


def discover_file(p):
    """-Path points at one image file. Returns (candidates, notes)."""
    p = Path(p)
    notes = []
    low = p.name.lower()
    if low.endswith(".vmdk"):
        cands, _n = discover(p.parent, max_depth=0)
        for c in cands:
            names = [Path(x).name.lower() for x in c["chain"]]
            if c.get("chain_desc") and low in names:
                i = names.index(low)
                descs = {Path(x).name.lower(): x for x in c["chain"][i:]}
                extents = [e for e in c["extents"] if e["descriptor"].lower() in descs]
                c = dict(c, leaf=c["chain"][i], chain=c["chain"][i:], chain_desc=c["chain_desc"][i:], extents=extents,
                         problems=[f"{e['extent']}: {e['status']}" for e in extents if e["status"] != "ok"])
                n = len(c["chain"])
                c["format"] = "VMDK" if n == 1 else f"VMDK (snapshot chain of {n} disks)"
                return _assign_labels([c], single_file=True), notes
        for q in sorted(p.parent.glob("*.vmdk")):
            d = read_descriptor(q) if q != p else None
            if d and any(e["file"].lower() == low for e in d["extents"]):
                notes.append(f"{p.name} is a data file of {q.name}; give the path of {q.name} instead")
                return [], notes
        if read_descriptor(p) is None:
            notes.append(f"{p.name} is not a VMDK descriptor; give the path of the descriptor .vmdk instead")
            return [], notes
        notes.append(f"{p.name}: could not follow its snapshot chain; opening it on its own")
        return _assign_labels([dict(image_cand(p, "VMDK", [p], []), segments=[])], single_file=True), notes
    if low.endswith(".ex01"):
        notes.append(f"Not opened (EWF2 .Ex01 is not supported; convert it to E01 or raw first): {p}")
        return [], notes
    if low.endswith(ARCHIVE_EXTS):
        notes.append(f"Not opened (archive; extract it first): {p}")
        return [], notes
    is_ewf = bool(_EWF_NUMBERED.match(p.suffix))
    if not is_ewf and _EWF_SEGMENT.match(p.suffix) and p.suffix.lower() not in RAW_IMAGE_EXTS:
        try:
            with open(lp(p), "rb") as fh:
                is_ewf = fh.read(3) == b"EVF"
        except OSError:
            is_ewf = False
    if is_ewf:
        first = p.with_suffix(p.suffix[:2] + "01") if _EWF_NUMBERED.match(p.suffix) else p
        if first != p and first.is_file():
            notes.append(f"{p.name} is part of an E01 set; opening its first segment {first.name}")
            p = first
        parts, problems = ewf_parts(p)
        if problems:
            notes.append(f"Not opened ({'; '.join(problems)}): {p}")
            return [], notes
        return _assign_labels([image_cand(p, "E01", parts, [])], single_file=True), notes
    if p.suffix[1:].isdigit():
        parts, problems = split_parts(p)
        if len(parts) > 1:
            if problems:
                notes.append(f"Not opened ({'; '.join(problems)}): {parts[0]}")
                return [], notes
            if parts[0] != p:
                notes.append(f"{p.name} is part of a split image; opening it from its first part {parts[0].name}")
            return _assign_labels([image_cand(parts[0], "split raw", parts, [])], single_file=True), notes
    fmt = OTHER_IMAGE_FORMATS.get(p.suffix.lower(), "raw")
    return _assign_labels([image_cand(p, fmt, [p], [])], single_file=True), notes


# ------------------------------------------------------------------ image helpers
def entry_row(p, query=""):
    """Metadata (+ SHA-256 for files) of one image path, as a dict."""
    r = {"path": str(p), "query": query, "type": "", "mode": "", "uid": "", "gid": "", "size": "", "mtime": "",
         "ctime": "", "btime": "", "atime": "", "sha256": "", "link": "", "note": ""}
    try:
        st = p.lstat()
    except Exception as ex:
        r["type"], r["note"] = "error", repr(ex)
        return r
    m = st.st_mode
    r["type"] = "dir" if stat.S_ISDIR(m) else "symlink" if stat.S_ISLNK(m) else "file" if stat.S_ISREG(m) else "other"
    r.update(mode=oct(m & 0o7777), uid=st.st_uid, gid=st.st_gid, size=st.st_size, mtime=ts(st.st_mtime),
             ctime=ts(st.st_ctime), btime=ts(getattr(st, "st_birthtime", None)), atime=ts(st.st_atime))
    if r["type"] == "symlink":
        try:
            r["link"] = str(p.readlink())
        except Exception as ex:
            r["note"] = repr(ex)
    elif r["type"] == "file":
        try:
            h = hashlib.sha256()
            with p.open("rb") as fh:
                for chunk in iter(lambda: fh.read(1 << 20), b""):
                    h.update(chunk)
            r["sha256"] = h.hexdigest()
        except Exception as ex:
            r["note"] = "hash failed: " + repr(ex)
    return r


def exists(t, path):
    try:
        p = t.fs.path(path)
        return p.exists() or p.is_symlink()
    except Exception:
        return False


def read_image_text(t, path, limit=1 << 20):
    try:
        with t.fs.path(path).open("rb") as fh:
            return fh.read(limit).decode("utf-8", "backslashreplace")
    except Exception:
        return None


def walk(t, top, limit=None):
    """Iterative walk that never follows symlinks. Yields (path, stat)."""
    stack, n = [t.fs.path(top)], 0
    while stack:
        d = stack.pop()
        try:
            entries = sorted(d.iterdir(), key=lambda x: x.name)
        except Exception:
            continue
        for e in entries:
            try:
                st = e.lstat()
            except Exception:
                continue
            n += 1
            if limit and n > limit:
                return
            if stat.S_ISDIR(st.st_mode):
                stack.append(e)
            yield e, st


def target_info(t):
    info = {}
    for key, fn in (("hostname", lambda: t.hostname), ("os", lambda: t.os), ("os_version", lambda: t.version),
                    ("ips", lambda: ", ".join(str(x) for x in (t.ips or []))),
                    ("install_date", lambda: ts(t.install_date)), ("last_activity", lambda: ts(t.activity))):
        try:
            v = fn()
            info[key] = "" if v is None else str(v)
        except Exception as ex:
            info[key] = f"(unavailable: {ex.__class__.__name__})"
    info["flash_version"] = (read_image_text(t, "/flash/.version") or "").strip()
    loader = read_image_text(t, "/flash/boot/loader.conf") or ""
    m = re.search(r'^\s*kernel\s*=\s*"?([^"\s]+)', loader, re.M)
    info["boot_kernel"] = m.group(1) if m else ""
    return info


# ------------------------------------------------------------------ file-system checks
def check_known_iocs(t):
    rows = []
    for path, desc in KNOWN_IOC_PATHS:
        if not exists(t, path):
            rows.append({"path": path, "query": path, "type": "missing", "mode": "", "uid": "", "gid": "", "size": "",
                         "mtime": "", "ctime": "", "btime": "", "atime": "", "sha256": "", "link": "", "note": "",
                         "description": desc, "present": False})
            continue
        p = t.fs.path(path)
        r = entry_row(p, path)
        r.update(description=desc, present=True)
        rows.append(r)
        if r["type"] == "dir":
            for e, _st in walk(t, path):
                c = entry_row(e, path)
                c.update(description=f"inside {path}", present=True)
                rows.append(c)
    return rows


# Boot scripts and crontabs. /flash/nsconfig/crontab is the persistent copy of /etc/crontab: Citrix (CTX327915) has
# administrators copy a customized /etc file into /nsconfig, and the appliance links it back into /etc at boot.
HOOK_FILES_GLOB = [("/flash/nsconfig", r"^(rc\.netscaler.*|nsafter\.sh|nsbefore\.sh|crontab)$"), ("/var/cron/tabs", r".*")]
SETUID_RX = re.compile(r"\bchmod\s+(?:-[A-Za-z]+\s+)*(?:[ugoa]*[+=][rwxXt]*s|0?[2-7][0-7]{3}\b)")
HTTPD_RELOAD_RX = re.compile(r"httpd\.pid|apachectl|\bhttpd\s+-k\b|kill\s+-(?:HUP|1|USR1)\b[^|;&]*httpd", re.I)
HTTPD_EDIT_RX = re.compile(r"\b(?:perl|sed|awk|ed|ex|tee|echo|printf)\b|>>?", re.I)
PHP_ENGINE_ON_RX = re.compile(r"php_(?:admin_)?flag\s+engine\s+on\b", re.I)


def check_hooks(t, since="", until=""):
    """Every line of the boot scripts and crontabs, with flags for what deserves a look. The files themselves are
    normal on a NetScaler (rc.netscaler holds ordinary settings), so only flagged lines are findings. A file created
    or modified in the review window gets a file-level row ("line" empty)."""
    rows = []
    x_rx = re.compile(IMPLANT_STRINGS, re.I)
    for folder, rx in HOOK_FILES_GLOB:
        if not exists(t, folder):
            continue
        for e in sorted(t.fs.path(folder).iterdir(), key=lambda x: x.name):
            if not re.match(rx, e.name):
                continue
            meta = entry_row(e)
            if meta["type"] != "file":
                continue
            text = read_image_text(t, str(e)) or ""
            base_flags = []
            if folder == "/var/cron/tabs":
                if e.name == "nobody":
                    base_flags.append("crontab owned by nobody")
                if meta["size"] == 0:
                    base_flags.append("empty crontab")
            changed = []
            if since and meta["btime"] and since <= meta["btime"] < until:
                changed.append("file created in the review window")
            if since and meta["mtime"] and since <= meta["mtime"] < until:
                changed.append("file modified in the review window")
            if not text.strip():
                rows.append({"file": str(e), "line": "", "text": "(empty file)", "flags": "; ".join(base_flags + changed),
                             "mtime": meta["mtime"], "btime": meta["btime"], "sha256": meta["sha256"], "installed": ""})
                continue
            if changed:
                rows.append({"file": str(e), "line": "", "text": f"(whole file: modified {meta['mtime']}, created "
                                                                  f"{meta['btime'] or 'unknown'})",
                             "flags": "; ".join(changed), "mtime": meta["mtime"], "btime": meta["btime"],
                             "sha256": meta["sha256"], "installed": ""})
            for i, line in enumerate(text.splitlines(), 1):
                s = line.strip()
                if not s:
                    continue
                flags = [] if s.startswith("#") else list(base_flags)
                installed = ""
                m = re.search(r"\((\S+) installed on (.+?)\)", s)
                if m:    # crontab header: flagged when the install time is in the review window or cannot be read
                    installed = f"installed from {m.group(1)} on {m.group(2)}"
                    when = cron_header_time(m.group(2))
                    if not when:
                        flags.append(installed)
                    elif since and since <= when < until:
                        flags.append(installed + " (in the review window)")
                if not s.startswith("#"):
                    if re.search(r"/\.[A-Za-z0-9_-]", s):
                        flags.append("runs something in a hidden folder")
                    if re.search(r"(^|[\s'\"])/(var/)?tmp/", s):
                        flags.append("runs from /tmp or /var/tmp")
                    if s.endswith("&"):
                        flags.append("starts a background process")
                    if re.search(r"\b(curl|wget|fetch|nc|ncat|socat|ssh|scp|base64)\b", s):
                        flags.append("network or encoding tool")
                    if IPV4_RX.search(s):
                        flags.append("contains an IP address")
                    if x_rx.search(s):
                        flags.append("names a known implant file or string")
                    if "httpd.conf" in s and HTTPD_EDIT_RX.search(s):
                        flags.append("edits the web server config (httpd.conf)")
                    if PHP_ENGINE_ON_RX.search(s):
                        flags.append("turns the PHP engine on")
                    if HTTPD_RELOAD_RX.search(s):
                        flags.append("reloads or restarts the web server")
                    if SETUID_RX.search(s):
                        flags.append("sets a setuid or setgid bit")
                rows.append({"file": str(e), "line": i, "text": line, "flags": "; ".join(flags),
                             "mtime": meta["mtime"], "btime": meta["btime"], "sha256": meta["sha256"],
                             "installed": installed})
    return rows


HTTPD_CONF_COPIES = ["/flash/nsconfig/httpd.conf"]   # persistent copy of /etc/httpd.conf, linked into /etc at boot
_BLOCK_OPEN_RX = re.compile(r"^<\s*(Directory|DirectoryMatch|Location|LocationMatch|Files|FilesMatch)\b\s*(.*?)\s*>$", re.I)
_BLOCK_CLOSE_RX = re.compile(r"^</\s*(Directory|DirectoryMatch|Location|LocationMatch|Files|FilesMatch)\s*>$", re.I)
_PHP_TYPE_RX = re.compile(r"^(AddType|AddHandler)\s+\S*php\S*\s+(.+)$", re.I)
_PHP_FORCE_RX = re.compile(r"^(ForceType|SetHandler)\s+\S*php", re.I)
_ALIAS_RX = re.compile(r"^(Alias|AliasMatch|ScriptAlias|ScriptAliasMatch|RewriteRule|Redirect\w*)\s+(.+)$", re.I)


def files_block_names_hidden_file(kind, arg):
    """True for <Files ".name"> or <FilesMatch "^\\.name"> naming a hidden file (not .ht* and not a bare dot rule)."""
    a = arg.strip()
    regex = kind.endswith("match") or a.startswith("~")
    a = a.lstrip("~").strip().strip("\"'")
    if regex:
        return bool(re.match(r"\^\\?\.(?!ht)[A-Za-z0-9]", a))
    return len(a) > 1 and a[0] == "." and a[1].isalnum() and not a.startswith(".ht")


def check_httpd_conf(t, since="", until=""):
    """Persistent httpd.conf copies: a summary row per file plus the lines that match what Unit 42 describes the actor
    adding (PHP engine switched on, aliases that expose a web shell as receiver.min.css, a Files block that runs a
    hidden file as PHP). Rows have the same shape as check_hooks rows."""
    rows = []
    x_rx = re.compile(IMPLANT_STRINGS, re.I)
    for path in HTTPD_CONF_COPIES:
        if not exists(t, path):
            continue
        e = t.fs.path(path)
        meta = entry_row(e)
        if meta["type"] != "file":
            continue
        text = read_image_text(t, path, limit=8 << 20) or ""
        lines = text.splitlines()
        has_off = bool(re.search(r"^\s*php_(?:admin_)?flag\s+engine\s+off\b", text, re.I | re.M))
        stack, found = [], []
        for i, line in enumerate(lines, 1):
            s = line.strip()
            if not s or s.startswith("#"):
                continue
            flags = []
            mo = _BLOCK_OPEN_RX.match(s)
            if mo:
                stack.append((mo.group(1).lower(), mo.group(2)))
                if mo.group(1).lower().startswith("files") and files_block_names_hidden_file(mo.group(1).lower(),
                                                                                             mo.group(2)):
                    flags.append("Files block for a hidden file (Unit 42: how the actor registered .ctxs.receiver)")
            elif _BLOCK_CLOSE_RX.match(s):
                if stack:
                    stack.pop()
            context = " ".join(arg for _k, arg in stack)
            if PHP_ENGINE_ON_RX.search(s):
                if line != line.rstrip():
                    flags.append("PHP engine switched on, with the trailing space the actor's one-line edit leaves "
                                 "(Unit 42)")
                elif not has_off:
                    flags.append("PHP engine switched on and no 'engine off' line left (NetScaler ships with it off; "
                                 "Unit 42)")
                elif re.search(r"logon|LogonPoint|custom", context, re.I):
                    flags.append("PHP engine switched on for the Gateway logon folder")
            ma = _ALIAS_RX.match(s)
            if ma:
                args = ma.group(2)
                target = args.split()[-1].strip("\"'")
                if re.search(r"receiver\.min|\.receiver\b", args, re.I) or re.search(r"/\.(?!ht|well-known)[A-Za-z0-9]",
                                                                                      target) \
                        or re.search(r"(?:^|\s)\"?/(?:var/)?tmp/", args):
                    flags.append("alias to a hidden file or a disguised web shell name (Unit 42: receiver.min.css and "
                                 "receiver.min.<hex>.css aliases)")
            mt_ = _PHP_TYPE_RX.match(s)
            if mt_:
                exts = [x for x in mt_.group(2).split() if x.startswith(".")]
                odd = [x for x in exts if x.lower() not in PHP_EXTS]
                if odd:
                    flags.append("runs PHP for non-PHP file types: " + " ".join(odd))
            if _PHP_FORCE_RX.match(s) and stack and "php" not in context.lower():
                flags.append(f"forces PHP inside <{stack[-1][0]} {stack[-1][1]}>")
            if x_rx.search(s):
                flags.append("names a known implant file or string")
            if flags:
                found.append({"file": path, "line": i, "text": line, "flags": "; ".join(dict.fromkeys(flags)),
                              "mtime": meta["mtime"], "btime": meta["btime"], "sha256": meta["sha256"], "installed": ""})
        changed = []
        if since and meta["btime"] and since <= meta["btime"] < until:
            changed.append("file created in the review window")
        if since and meta["mtime"] and since <= meta["mtime"] < until:
            changed.append("file modified in the review window")
        rows.append({"file": path, "line": "", "text": f"(persistent copy of /etc/httpd.conf, used at boot: {len(lines)} "
                                                       f"lines, {len(found)} flagged; only flagged lines are listed)",
                     "flags": "; ".join(changed), "mtime": meta["mtime"], "btime": meta["btime"],
                     "sha256": meta["sha256"], "installed": ""})
        rows.extend(found)
    return rows


CORE_DIRS = ("/var/core", "/var/crash")
CORE_SKIP = {"minfree", "bounds", ".snap"}      # FreeBSD bookkeeping files in /var/crash


def list_cores(t):
    """Core and crash dumps on the disk with their times. Unit 42 says the CVE-2026-88772 DTLS exploit causes memory
    corruption or a crash, and crashing processes (packet engines, nsaaad) can leave cores, so these times are worth
    lining up with the C and D crash records."""
    rows = []
    for top in CORE_DIRS:
        if not exists(t, top):
            continue
        for e, st in walk(t, top, limit=20000):
            if not stat.S_ISREG(st.st_mode) or e.name in CORE_SKIP:
                continue
            low = e.name.lower()
            kind = ("packet engine (NSPPE) core" if "nsppe" in low
                    else "kernel crash dump" if re.match(r"(vmcore|info|textdump|core\.txt)(\.|$)", low)
                    else "process core")
            rows.append({"path": str(e), "kind": kind, "size": st.st_size, "mtime": ts(st.st_mtime),
                         "btime": ts(getattr(st, "st_birthtime", None))})
    rows.sort(key=lambda r: (r["mtime"], r["path"]))
    return rows


def cron_header_time(text):
    """'Mon Dec  6 10:00:00 2021' (crontab header, appliance clock) -> '2021-12-06 10:00:00', or ''."""
    m = re.search(r"\w{3} (\w{3}) +(\d+) (\d\d:\d\d:\d\d) (\d{4})", text)
    if not m or m.group(1) not in MONTHS:
        return ""
    return f"{m.group(4)}-{MONTHS[m.group(1)]:02d}-{int(m.group(2)):02d} {m.group(3)}"


def read_head_bytes(e, n):
    try:
        with e.open("rb") as fh:
            return fh.read(n)
    except Exception:
        return None


def sha256_image_file(e):
    try:
        h = hashlib.sha256()
        with e.open("rb") as fh:
            for chunk in iter(lambda: fh.read(1 << 20), b""):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return ""


def file_checks(e, p, kind, size):
    """(Discovery category or None, SHA-256 or '') for one file: the VPN / EPA client download folder checks, PHP code
    hidden in a non-PHP file under a web folder, and the published hashes (KNOWN_SHA256)."""
    if kind != "file":
        return None, ""
    name = p.rsplit("/", 1)[1]
    ext = os.path.splitext(name)[1].lower()
    in_pkg = bool(CLIENT_PKG_DIR_RX.match(p))
    cat = None
    if in_pkg and name.lower() in UNIT42_PKG_NAMES:
        cat = "Lead: client download folder: file name Unit 42 reported as a webshell"
    head = None
    if WEB_ROOT_RX.match(p) and ext not in PHP_EXTS and (
            not GUI_ROOT_RX.match(p) or in_pkg or ext in STATIC_EXTS or not ext or name.startswith(".")):
        head = read_head_bytes(e, 65536)
        if head is not None and PHP_CODE_RX.search(head):
            cat = ("Lead: client download folder: PHP code in a package-named file, a disguised web shell" if in_pkg
                   else "Lead: web folder: PHP code in a non-PHP file (possible web shell)")
    if cat is None and in_pkg and ext in PKG_MAGIC:
        head = head if head is not None else read_head_bytes(e, 8)
        if head is None:
            cat = "Lead: client download folder: package could not be read"
        elif not head.startswith(PKG_MAGIC[ext]):
            cat = "Lead: client download folder: not a real package, possibly a disguised webshell"
    sha = ""
    if (cat or HASH_SCOPE_RX.match(p)) and size <= HASH_MAX_BYTES:
        sha = sha256_image_file(e)
        if sha in KNOWN_SHA256:
            cat = f"IOC: SHA-256 matches {KNOWN_SHA256[sha]}"
    return cat, sha


def discovery_walk(t, since, until):
    """Hidden, implant-like, and new-in-window items in /flash and /var, with a category for each."""
    rows = []
    s, u = ts(since), ts(until)
    for top in ("/flash", "/var"):
        for e, st in walk(t, top):
            p = str(e)
            m = st.st_mode
            kind = "dir" if stat.S_ISDIR(m) else "symlink" if stat.S_ISLNK(m) else "file"
            b = ts(getattr(st, "st_birthtime", None))
            mt = ts(st.st_mtime)
            hidden = any(x.startswith(".") for x in p.split("/") if x)
            born_in = bool(b) and s <= b < u
            fcat, sha = file_checks(e, p, kind, st.st_size)
            cat = None
            if fcat and fcat.startswith("IOC"):
                cat = fcat
            elif KNOWN_IOC_RX.search(p):
                cat = "IOC: known implant artifact"
            elif fcat:
                cat = fcat
            elif p == LOGON_DROP_DIR and kind == "dir" and born_in:
                cat = "Lead: Gateway logon folder custom/ created in the window (Unit 42: the actor's web shell drop path)"
            elif hidden and BENIGN_HIDDEN_RX.search(p):
                cat = "Context: standard hidden file"
            elif hidden:
                cat = ("Lead: non-standard hidden item created in the window" if born_in
                       else "Context: non-standard hidden item older than the window")
            elif born_in and not p.startswith(WALK_SKIP_WINDOW):
                if kind == "file" and st.st_uid == 65534 and not WEB_NOISE_RX.search(p) and not re.search(r"/sess_[0-9a-f]+$", p):
                    cat = "Lead: written by the web server (nobody) in the window"
                elif p.startswith("/flash/") and not FLASH_ROTATING_RX.search(p) and kind != "dir":
                    cat = "Lead: new file in /flash in the window"
                elif re.match(r"^/var/(tmp|nstmp)/", p) and not TMP_NOISE_RX.search(p) and not re.search(r"/sess_[0-9a-f]+$", p):
                    cat = "Lead: new item in /var/tmp or /var/nstmp in the window"
                elif re.search(r"/sess_[0-9a-f]+$", p):
                    cat = "Context: GUI session file in the window"
                else:
                    cat = "Context: other item created in the window"
            if cat:
                rows.append({"category": cat, "path": p, "type": kind, "mode": oct(m & 0o7777), "uid": st.st_uid,
                             "size": st.st_size, "mtime": mt, "ctime": ts(st.st_ctime), "btime": b, "sha256": sha})
    return rows


# ------------------------------------------------------------------ export
def export_tree(t, img_root, dest_root, manifest, group):
    """Copies regular files under img_root (recursively) to dest_root; one manifest row per entry."""
    if not exists(t, img_root):
        return 0
    n = 0
    base = img_root.rstrip("/")
    root = t.fs.path(img_root)
    is_dir = root.is_dir() and not root.is_symlink()
    entries = walk(t, img_root) if is_dir else [(root, root.lstat())]
    for e, st in entries:
        p = str(e)
        if stat.S_ISDIR(st.st_mode):
            continue
        rel = p[len(base):].lstrip("/") if is_dir else Path(p).name
        out = Path(dest_root) / safe_rel(rel)
        row = {"group": group, "image_path": p, "exported_to": "", "size": st.st_size, "mode": oct(st.st_mode & 0o7777),
               "uid": st.st_uid, "mtime": ts(st.st_mtime), "ctime": ts(st.st_ctime),
               "btime": ts(getattr(st, "st_birthtime", None)), "sha256": "", "note": ""}
        if stat.S_ISLNK(st.st_mode):
            try:
                row["note"] = "symlink to " + str(e.readlink())
            except Exception:
                row["note"] = "symlink"
            manifest.append(row)
            continue
        if not stat.S_ISREG(st.st_mode):
            row["note"] = "not a regular file"
            manifest.append(row)
            continue
        k = 2
        while os.path.exists(lp(out)):
            out = out.with_name(f"{out.stem}_{k}{out.suffix}")
            k += 1
        try:
            os.makedirs(lp(out.parent), exist_ok=True)
            h = hashlib.sha256()
            with e.open("rb") as src, open(lp(out), "wb") as dst:
                for chunk in iter(lambda: src.read(1 << 20), b""):
                    h.update(chunk)
                    dst.write(chunk)
            try:
                os.utime(lp(out), (st.st_atime, st.st_mtime))
            except Exception:
                pass
            row["sha256"], row["exported_to"] = h.hexdigest(), str(out)
            n += 1
        except Exception as ex:
            row["note"] = "export failed: " + repr(ex)
        manifest.append(row)
    return n


# ------------------------------------------------------------------ log pass: xlsx conversion + hunts + integrity
def parse_utx(path, fixed):
    """FreeBSD utmpx records (utx.log*: 2-byte big-endian length prefix; utx.active/lastlogin: fixed 197 bytes)."""
    data = b""
    opener = gzip.open if str(path).lower().endswith(".gz") else open
    with opener(lp(path), "rb") as fh:
        data = fh.read()
    recs, off = [], 0
    while off < len(data):
        if fixed:
            raw = data[off:off + 197]
            off += 197
            if len(raw) < 197:
                break
        else:
            if off + 2 > len(data):
                break
            (n,) = struct.unpack(">H", data[off:off + 2])
            off += 2
            if n == 0 or off + n > len(data):
                break
            raw = data[off:off + n] + b"\x00" * max(0, 197 - n)
            off += n
        typ = raw[0]
        tv = struct.unpack(">Q", raw[1:9])[0]

        def s(b):
            return b.split(b"\x00")[0].decode("utf-8", "backslashreplace")
        recs.append({"type": UTX_TYPES.get(typ, str(typ)), "time_utc": ts(tv / 1e6) if tv else "",
                     "pid": struct.unpack(">I", raw[17:21])[0], "user": s(raw[21:53]), "line": s(raw[53:69]),
                     "host": s(raw[69:197]), "id": raw[9:17].hex()})
    return recs


def norm_ip(text):
    """'203.0.113.5' or an IPv6 address in any spelling -> canonical text; None when it is not an IP address."""
    try:
        return ipaddress.ip_address(str(text).strip().strip("[]")).compressed
    except ValueError:
        return None


def ioc_ip_table(extra):
    """{ip: (level, note)}: the built-in list plus the IPs given with --ioc-ip (built-in notes win)."""
    table = {norm_ip(ip): (level, note) for ip, level, note in BUILTIN_IOC_IPS}
    for ip in extra:
        table.setdefault(norm_ip(ip), SUPPLIED_IOC)
    return table


def ip_regex(ips):
    """One regex that finds any of the IPs as a whole address (not inside a longer number); group 1 is the IP."""
    pats = []
    for ip in sorted(ips, key=len, reverse=True):
        if ":" in ip:
            pats.append(r"(?<![0-9A-Fa-f:])" + re.escape(ip) + r"(?![0-9A-Fa-f:])")
        else:
            pats.append(r"(?<![\d.])" + re.escape(ip) + r"(?![\d.])")
    return re.compile("(" + "|".join(pats) + ")", re.I) if pats else None


class IpMatcher:
    """Finds any of the attacker IPs in a line as a whole address. IPv4 candidates are looked up in a set, so hundreds
    of IPs cost no more than a few; IPv6 IOCs use a regex. search() returns a match whose group(1) is the IP, or None."""

    def __init__(self, ips):
        self.v4 = {ip for ip in ips if ":" not in ip}
        self.v6 = ip_regex([ip for ip in ips if ":" in ip])

    def search(self, text):
        if self.v4:
            for m in IPV4_RX.finditer(text):
                if m.group(1) in self.v4:
                    return m
        return self.v6.search(text) if self.v6 is not None else None


class Hunts:
    def __init__(self, ioc_ips):
        self.ioc_ips = dict(ioc_ips)          # {ip: (level, note)}
        self.ip_rx = IpMatcher(self.ioc_ips) if self.ioc_ips else None
        self.sections = []
        for key, sheet, title, pats, rx in SECTIONS:
            crx = self.ip_rx if key == "H" else re.compile(rx, re.I) if isinstance(rx, str) else rx
            if crx is None:
                continue
            self.sections.append((key, sheet, title, pats, crx))
        self.hits = {k: [] for k, *_ in self.sections}
        self.e_rx = re.compile(next(s[4] for s in SECTIONS if s[0] == "E"), re.I)
        self.forged = PitbossMatcher()
        self.x_rx = re.compile(IMPLANT_STRINGS, re.I)
        self.x = {}
        self.nsppe = []      # (file, line, time, text, section) of every A and B hit: the PITBOSS injections
        self.audit = {}
        self.nul = []
        self.utx = []
        self.clock = []      # (syslog prefix time, GMT time) pairs from ns.log for the clock check

    def ip_info(self, ip_text):
        """(level, note) of a matched attacker IP."""
        return self.ioc_ips.get(norm_ip(ip_text), ("", ""))

    def applicable(self, name, top_level):
        name = name.lower()
        out = []
        for k, _s, _t, pats, rx in self.sections:
            if not (top_level or k == "H"):
                continue
            if any(fnmatch.fnmatchcase(name, p) for p in pats if not p.startswith("!")) and \
                    not any(fnmatch.fnmatchcase(name, p[1:]) for p in pats if p.startswith("!")):
                out.append((k, rx))
        return out

    def line(self, rel, ln, text, t_utc, act, top_level, file_name):
        forged = None
        for key, rx in act:
            m = rx.search(text)
            if not m:
                continue
            if key == "E2" and self.e_rx.search(text):
                continue
            if key == "I" and SECTION_I_EXCLUDE.search(text):
                continue
            if key in ("C", "D"):            # genuine crash records only; forged ones are B
                if forged is None:
                    forged = bool(self.forged.search(text))
                if forged:
                    continue
            self.hits[key].append((rel, ln, t_utc, m.group(1) if key == "H" else "", text))
            if key in ("A", "B"):
                self.nsppe.append((rel, ln, t_utc, text, key))
        m = self.x_rx.search(text)
        if m:
            e = self.x.setdefault(rel, [0, "", "", set()])
            e[0] += 1
            e[3].add(m.group(0).lower())
            if t_utc:
                e[1] = e[1] or t_utc
                e[2] = t_utc
        if top_level:
            low = file_name.lower()
            if low.startswith("notice.log") and 'sh_command="' in text:
                cmd = text.split('sh_command="', 1)[1].rstrip().rstrip('"').strip()
                cmd = re.sub(r"[A-Za-z0-9+/=]{40,}", "<B64>", cmd)
                e = self.audit.setdefault(cmd, [0, "", "", rel])
                e[0] += 1
                if t_utc:
                    e[1] = e[1] or t_utc
                    e[2] = t_utc
            if low.startswith("ns.log") and len(self.clock) < 50:
                m1, m2 = SYSLOG_TS.match(text), NSGMT_TS.search(text)
                if m1 and m2:
                    self.clock.append((f"{m1.group(3)}:{m1.group(4)}:{m1.group(5)}", f"{m2.group(4)}:{m2.group(5)}:{m2.group(6)}"))


def process_log(local, rel, mtime_s, top_level, xlsx_out, hunts, skip_xlsx):
    """One pass over an exported log: classify, convert to .xlsx, feed the hunts, record NUL runs."""
    info = {"kind": "", "lines": 0, "nul_lines": 0, "nul_bytes": 0, "xlsx": "", "note": ""}
    name = Path(rel).name
    ref = parse_ts(mtime_s)
    low = name.lower()
    try:
        if UTX_LOG_RX.match(low) or UTX_FIXED_RX.match(low):
            recs = parse_utx(local, fixed=bool(UTX_FIXED_RX.match(low)))
            info.update(kind="binary: utx login records", lines=len(recs))
            for r in recs:
                hunts.utx.append((rel, r))
                if hunts.ip_rx is not None and r["host"]:
                    m = hunts.ip_rx.search(r["host"])
                    if m:
                        hunts.hits["H"].append((rel, "", r["time_utc"], m.group(1),
                                                f"utx {r['type']} user={r['user']} line={r['line']} host={r['host']}"))
            if not skip_xlsx and recs:
                b = Book(xlsx_out)
                b.table("records", ["Type", "Time (UTC)", "User", "Line (tty)", "Host", "PID", "ID (hex)"],
                        ([r["type"], r["time_utc"], r["user"], r["line"], r["host"], r["pid"], r["id"]] for r in recs),
                        [16, 20, 16, 12, 30, 10, 18])
                b.close()
                info["xlsx"] = str(xlsx_out)
            return info
        head = read_head(local)
        if not head:
            info["kind"] = "empty"
            return info
        if looks_binary(head):
            info["kind"] = "binary (not converted)"
            if top_level and hunts.ip_rx is not None:     # same as the PowerShell hunt: binary files are searched too
                for ln, raw in enumerate(iter_lines(local), 1):
                    text = raw.decode("latin-1")
                    m = hunts.ip_rx.search(text)
                    if m:
                        hunts.hits["H"].append((rel, ln, "", m.group(1), "(binary file) " + text[:300]))
            return info
        info["kind"] = "text"
        act = hunts.applicable(name, top_level)
        is_http = low.startswith("httpaccess")
        book = None
        if not skip_xlsx:
            book = Book(xlsx_out)
            if is_http:
                headers = ["Line", "Time (UTC)", "Client IP", "Via", "Method", "Path", "Query", "Status", "Bytes",
                           "User agent", "Text"]
                widths = [8, 20, 16, 14, 8, 50, 30, 7, 9, 40, 120]
            else:
                headers, widths = ["Line", "Time (UTC)", "Text"], [8, 20, 160]
        pending_nul = []
        last_time = ""

        def row_gen():
            nonlocal last_time
            for ln, raw in enumerate(iter_lines(local), 1):
                info["lines"] = ln
                nul = raw.count(b"\x00")
                text = raw.decode("utf-8", "backslashreplace").rstrip("\r\n")
                clean = text.lstrip("\x00")
                t_utc = line_time(clean, ref)
                if nul:
                    info["nul_lines"] += 1
                    info["nul_bytes"] += nul
                    pending_nul.append([rel, ln, nul, last_time, t_utc if clean.strip("\x00") else ""])
                elif pending_nul and t_utc:
                    for pn in pending_nul:
                        if not pn[4]:
                            pn[4] = t_utc
                    hunts.nul.extend(pending_nul)
                    pending_nul.clear()
                if t_utc:
                    last_time = t_utc
                hunts.line(rel, ln, text, t_utc, act, top_level, name)
                if book is not None:
                    if is_http:
                        m = HTTPACC_RX.match(clean)
                        if m:
                            req = (m.group("req") or "").split(" ")
                            method = req[0] if req else ""
                            target = req[1] if len(req) > 1 else ""
                            path, _, query = target.partition("?")
                            yield [ln, t_utc, m.group("ip"), m.group("via") or "", method, path, query, m.group("status"),
                                   m.group("bytes"), m.group("ua") or "", text]
                        else:
                            yield [ln, t_utc, "", "", "", "", "", "", "", "", text]
                    else:
                        yield [ln, t_utc, text]

        if book is not None:
            book.table("log", headers, row_gen(), widths)
            book.close()
            info["xlsx"] = str(xlsx_out)
        else:
            for _ in row_gen():
                pass
        if pending_nul:
            hunts.nul.extend(pending_nul)
        if info["nul_lines"]:
            info["note"] = f"{info['nul_bytes']} NUL bytes on {info['nul_lines']} line(s)"
    except (EOFError, OSError, gzip.BadGzipFile) as ex:
        info["note"] = (info["note"] + "; " if info["note"] else "") + f"read stopped early: {ex.__class__.__name__}: {ex}"
    return info


# ------------------------------------------------------------------ derived analyses
def summarize_sections(hunts):
    out = []
    for key, sheet, title, pats, _rx in hunts.sections:
        rows = hunts.hits[key]
        times = sorted(r[2] for r in rows if r[2])
        files = collections.Counter(r[0] for r in rows)
        inc = ", ".join(p for p in pats if not p.startswith("!"))
        exc = ", ".join(p[1:] for p in pats if p.startswith("!"))
        out.append({"key": key, "sheet": sheet, "title": title, "files_filter": inc + (f" (except {exc})" if exc else ""),
                    "count": len(rows),
                    "count_top_level": sum(1 for r in rows if "/" not in r[0]),
                    "first": times[0] if times else "", "last": times[-1] if times else "",
                    "files": "; ".join(f"{f} ({n})" for f, n in files.most_common())})
    return out


def webshell_rows(hunts):
    rows = []
    for rel, ln, t_utc, _d, text in hunts.hits.get("E", []):
        m = HTTPACC_RX.match(text.lstrip("\x00"))
        if not m:
            rows.append({"time": t_utc, "ip": "", "method": "", "path": "", "status": "", "bytes": "", "ua": "",
                         "file": rel, "line": ln})
            continue
        req = (m.group("req") or "").split(" ")
        rows.append({"time": t_utc, "ip": m.group("ip"), "method": req[0] if req else "",
                     "path": (req[1] if len(req) > 1 else "").split("?")[0], "status": m.group("status"),
                     "bytes": m.group("bytes"), "ua": m.group("ua") or "", "file": rel, "line": ln})
    return rows


def shell_distinct(hunts):
    """Section I lines reduced to distinct commands (process ids and long base64 runs removed)."""
    g = {}
    for rel, ln, t_utc, _d, text in hunts.hits.get("I", []):
        if 'sh_command="' in text:
            cmd = text.split('sh_command="', 1)[1].rstrip().rstrip('"')
        else:
            m = re.search(r"\]:\s(.*)$", text)
            cmd = m.group(1) if m else text
        cmd = re.sub(r"[A-Za-z0-9+/=]{40,}", "<B64>", re.sub(r"\s+", " ", cmd)).strip()[:300]
        e = g.setdefault(cmd, {"n": 0, "first": "", "last": "", "files": collections.Counter()})
        e["n"] += 1
        e["files"][rel] += 1
        if t_utc:
            e["first"] = min(e["first"], t_utc) if e["first"] else t_utc
            e["last"] = max(e["last"], t_utc)
    out = [{"command": k, "count": v["n"], "first": v["first"], "last": v["last"],
            "files": "; ".join(f"{f} ({n})" for f, n in v["files"].most_common(8))} for k, v in g.items()]
    out.sort(key=lambda x: (x["count"], x["first"]))
    return out


def webshell_by_ip(ws_rows):
    g = collections.OrderedDict()
    for r in sorted(ws_rows, key=lambda x: x["time"]):
        e = g.setdefault((r["ip"], r["path"]), {"n": 0, "first": "", "last": "", "status": collections.Counter(),
                                                 "with_body": 0, "days": set()})
        e["n"] += 1
        e["first"] = e["first"] or r["time"]
        e["last"] = r["time"] or e["last"]
        e["status"][r["status"]] += 1
        if r["bytes"].isdigit() and int(r["bytes"]) > 0:
            e["with_body"] += 1
        if r["time"]:
            e["days"].add(r["time"][:10])
    out = []
    for (ip, path), e in g.items():
        out.append({"ip": ip, "path": path, "count": e["n"], "with_body": e["with_body"], "first": e["first"],
                    "last": e["last"], "days": ", ".join(sorted(e["days"])),
                    "status": ", ".join(f"{k}x{v}" for k, v in e["status"].most_common())})
    out.sort(key=lambda x: -x["count"])
    return out


NO_COMMAND = "(no command: pitboss text in a user name only)"


def injected_command(text):
    """The command carried by a forged pitboss / NSPPE record (URL-decoded), or NO_COMMAND."""
    if "%" in text:
        text = urllib.parse.unquote(text)
    m = INJECT_CMD_RX.search(text)
    cmd = m.group("cmd") if m else ""
    if not cmd.strip():
        m = INJECT_SUBST_RX.search(text)
        if m:
            cmd = m.group(1) or m.group(2) or ""
    cmd = re.sub(r"\s+", " ", cmd).strip()
    return cmd[:300] if cmd else NO_COMMAND


def nsppe_summary(hunts):
    """PITBOSS injections (hunts A and B) grouped by injected command."""
    found = collections.OrderedDict()
    for rel, ln, t_utc, text, key in hunts.nsppe:
        cmd = injected_command(text)
        e = found.setdefault(cmd, {"n": 0, "by": collections.Counter(), "times": [], "ips": set(),
                                   "files": collections.Counter()})
        e["n"] += 1
        e["by"][key] += 1
        if t_utc:
            e["times"].append(t_utc)
        for a, b in CLIENT_IP_RX.findall(text):
            ip = norm_ip(a or b)
            if ip:
                e["ips"].add(ip)
        e["files"][rel] += 1
    out = []
    for cmd, e in found.items():
        urls = sorted(set(URL_RX.findall(cmd)))
        out.append({"injected": cmd, "count": e["n"], "forged": e["by"]["B"], "callhome": e["by"]["A"],
                    "first": min(e["times"]) if e["times"] else "", "last": max(e["times"]) if e["times"] else "",
                    "client_ips": ", ".join(sorted(e["ips"])),
                    "files": "; ".join(f"{f} ({n})" for f, n in e["files"].most_common()),
                    "urls": " ".join(urls), "ips_in_text": ", ".join(sorted(set(IPV4_RX.findall(cmd))))})
    out.sort(key=lambda x: (x["injected"] == NO_COMMAND, x["first"] or "9999"))
    return out


def ua_payload_rows(hunts):
    """Hunt B2 hits with the payload measured and hashed. The decoded content itself is never written out, only
    its kind, SHA-256 and the URLs and IPs inside it."""
    rows = []
    for rel, ln, t_utc, _d, text in hunts.hits.get("B2", []):
        clean = text.lstrip("\x00")
        m = HTTPACC_RX.match(clean)
        ua = user_agent(clean)
        mark, b64 = ua_payload(ua) or ("", "")
        raw = decode_b64(b64)
        dec = raw.decode("utf-8", "replace")
        hashes = {hashlib.sha256(b64.encode()).hexdigest(), hashlib.sha256((b64 + "\n").encode()).hexdigest()}
        h_dec = hashlib.sha256(raw).hexdigest() if raw else ""
        known = sorted({KNOWN_SHA256[h] for h in hashes | {h_dec} if h in KNOWN_SHA256})
        req = (m.group("req") or "").split(" ") if m else []
        ip = m.group("ip") if m else ""
        if not ip:
            mi = IPV4_RX.search(clean)
            ip = mi.group(1) if mi else ""
        rows.append({"time": t_utc, "ip": ip, "method": req[0] if req else "",
                     "path": (req[1] if len(req) > 1 else "").split("?")[0], "status": m.group("status") if m else "",
                     "marker": mark, "length": len(b64), "kind": payload_kind(raw),
                     "sha_b64": hashlib.sha256(b64.encode()).hexdigest(), "sha_dec": h_dec, "known": "; ".join(known),
                     "urls": " ".join(sorted(set(URL_RX.findall(dec))))[:2000],
                     "ips": ", ".join(sorted(set(IPV4_RX.findall(dec))))[:1000], "ua": ua, "file": rel, "line": ln})
    return rows


def agent_log_rows(local, mtime_s):
    """Dates the [HH:MM:SS] lines of agent.log by counting day changes backwards from the file's modified time."""
    lines = []
    with open(lp(local), encoding="utf-8", errors="backslashreplace") as fh:
        for raw in fh:
            m = re.match(r"^\[(\d\d):(\d\d):(\d\d)\] (.*)", raw.rstrip("\r\n"))
            if m:
                lines.append((dt.time(int(m.group(1)), int(m.group(2)), int(m.group(3))), m.group(4)))
    end = parse_ts(mtime_s)
    if not lines or end is None:
        return []
    day = end.date()
    # Going backwards, a line more than an hour "later" in the day than the line after it crossed midnight. The file's
    # modified time counts as the line after the last one, so a last line later in the day than it is from the day before.
    out, nxt = [], end.time().replace(microsecond=0)
    for t_, msg in reversed(lines):
        if t_ > nxt and (dt.datetime.combine(day, t_) - dt.datetime.combine(day, nxt)).seconds > 3600:
            day -= dt.timedelta(days=1)
        nxt = t_
        kind = "agent start" if msg.startswith("agent listening") else "open" if msg.startswith("open ") else (
            "close" if "close" in msg.lower() else "other")
        m = re.match(r"open (\S+?):(\d+)(?: -> sid (\d+))?", msg)
        out.append({"date": day.isoformat(), "time": t_.strftime("%H:%M:%S"), "kind": kind, "message": msg,
                    "dest_ip": m.group(1) if m else "", "dest_port": m.group(2) if m else "",
                    "sid": (m.group(3) or "") if m else ""})
    out.reverse()
    return out


def timeline(label, ioc_rows, hooks, sec_summary, ws_by_ip, hunts, agent_rows, logging_stopped, vmsd, since_s, cores=()):
    ev = []

    def add(t_, src, event, detail=""):
        if t_:
            ev.append({"time": t_, "target": label, "source": src, "event": event, "detail": detail})
    for c_ in cores:
        add(c_["mtime"], "core dumps", f"{c_['kind']} written: {c_['path']}", f"{c_['size']:,} bytes (file modified time)")
    for r in ioc_rows:
        if r.get("present") and r["type"] in ("file", "dir"):
            add(r["btime"], "file system", f"created {r['path']}", f"{r['type']}, {r['size']} bytes, sha256 {r['sha256'][:16]}")
            if r["mtime"] and r["mtime"] != r["btime"]:
                add(r["mtime"], "file system", f"modified {r['path']}", f"size now {r['size']} bytes")
    for h in hooks:
        m = re.search(r"installed from (\S+) on (.+)$", h.get("installed", ""))
        if m and cron_header_time(m.group(2)):
            add(cron_header_time(m.group(2)), "crontab header", f"crontab {Path(h['file']).name} installed from {m.group(1)}",
                "appliance clock")
    for s in sec_summary:
        if s["count"]:
            add(s["first"], f"log hunt {s['key']}", f"first: {s['title']}", f"{s['count']} line(s) in total")
            if s["last"] != s["first"]:
                add(s["last"], f"log hunt {s['key']}", f"last: {s['title']}", "")
    for w in ws_by_ip[:40]:
        add(w["first"], "httpaccess", f"first webshell request from {w['ip']} to {w['path']}",
            f"{w['count']} request(s), {w['with_body']} with a response body, until {w['last']}")
    for rel, ln, t_utc, _d, text in hunts.hits.get("D", []):
        add(t_utc, f"/var/log/{rel}", "packet engine (NSPPE) crash", text[-120:])
    logins = collections.OrderedDict()
    for rel, r in sorted(hunts.utx, key=lambda x: x[1]["time_utc"]):
        if r["type"] == "USER_PROCESS" and r["time_utc"] >= since_s:
            e = logins.setdefault((r["user"], r["host"] or "(local)", r["time_utc"][:10]), [0, r["time_utc"], ""])
            e[0] += 1
            e[2] = r["time_utc"]
    for (user, host, _day), (n, first, last) in logins.items():
        add(first, "utx.log", f"{n} login(s) by {user} from {host}", f"first {first}, last {last}")
    for r in agent_rows:
        if r["kind"] == "open":
            add(f"{r['date']} {r['time']}", "/var/tmp/.ux/agent.log", f"tunnel open to {r['dest_ip']}:{r['dest_port']}",
                "date inferred from the file's modified time")
    add(logging_stopped, "/var/log", "newest log modification (logging stops)", "")
    for s in vmsd:
        add(s["created_utc"], ".vmsd", f"VM snapshot created: {s['name']}", "")
    ev.sort(key=lambda x: (x["time"], x["event"]))
    seen, out = set(), []
    for e in ev:
        k = (e["time"], e["event"])
        if k not in seen:
            seen.add(k)
            out.append(e)
    return out


# ------------------------------------------------------------------ Fox-IT iocitrix
def run_foxit(t, foxit_path, out_txt):
    spec = importlib.util.spec_from_file_location("iocitrix", foxit_path)
    ioc = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ioc)
    FindingRecord = ioc.FindingRecord

    def owner(record):
        user = getattr(record, "username", None)
        if user:
            return user
        source = str(record.source).replace("\\", "/") if getattr(record, "source", None) else ""
        return source.rsplit("/", 1)[-1] if source.startswith("/var/cron/tabs/") else None

    def check_crontabs(target):
        for record in target.cronjobs():
            if record._desc.name.endswith("environmentvariable"):
                continue
            if owner(record) in ["nobody"]:
                yield FindingRecord(type="cronjob/user", alert="Crontab by nobody user observed", confidence="high",
                                    path=record.source)
            for name, pattern in ioc.EVIL_CRONTAB_CONTENTS:
                match = pattern.search(record.command or "")
                if match:
                    yield FindingRecord(type="cronjob/command", confidence="medium", path=record.source,
                                        alert=f"{name.strip()} found in crontab command ({match.group(0)})")

    def check_timestomps(target):
        for d in ioc.TIMESTOMP_DIRS:
            base = target.fs.path(d)
            if not base.exists():
                continue
            for entry in base.rglob("*"):
                if not entry.exists():
                    continue
                st = entry.lstat()
                seconds = int(st.st_ctime - st.st_mtime)
                if seconds > ioc.TIMESTOMP_THRESHOLD_SECONDS:
                    yield FindingRecord(type="file/timestomp", confidence="medium", path=entry,
                                        alert=f"Possibly Timestomped File Observed (ctime is {seconds // 86400} days after mtime)")

    ioc.check_crontabs = check_crontabs
    ioc.check_timestomps = check_timestomps
    with open(lp(out_txt), "w", encoding="utf-8") as fh, contextlib.redirect_stdout(fh):
        findings = ioc.ioc_check_target(t)
    return [{"confidence": f.confidence, "type": f.type, "alert": f.alert, "path": str(f.path)} for f in findings]


# ------------------------------------------------------------------ per-image hunt
def hunt_image(cand, img_path, label, out_root, args, run):
    from dissect.target import Target
    tdir = Path(out_root) / label
    os.makedirs(lp(tdir), exist_ok=True)
    res = {"label": label, "image": img_path, "cand": cand, "status": "", "errors": []}
    say(f"[{label}] opening {img_path}")
    t = Target.open(img_path)
    info = target_info(t)
    res["info"] = info
    if info.get("os") != "citrix-netscaler":
        res["status"] = f"not a NetScaler root disk (os={info.get('os')}); inventoried only"
        return res | inventory_disk(t, label, tdir, args)
    say(f"[{label}] {info.get('hostname')} {info.get('os_version')} (last activity {info.get('last_activity')})")

    # window
    manifest = []
    say(f"[{label}] exporting /var/log and the implant text logs")
    logs_dir = tdir / "logs"
    export_tree(t, "/var/log", logs_dir / "var_log", manifest, "var_log")
    for p in IMPLANT_LOG_EXPORTS:
        if exists(t, p):
            export_tree(t, p, logs_dir / "implant_logs" / safe_rel(p.rsplit("/", 1)[0]), manifest, "implant_logs")
    var_log_times = sorted(m["mtime"] for m in manifest if m["group"] == "var_log" and m["mtime"])
    logging_stopped = var_log_times[-1] if var_log_times else ""
    anchor = parse_ts(logging_stopped) or parse_ts(info.get("last_activity", "")) or dt.datetime.now(UTC)
    since = (dt.datetime.strptime(args.since, "%Y-%m-%d").replace(tzinfo=UTC) if args.since
             else (anchor - dt.timedelta(days=30)).replace(hour=0, minute=0, second=0))
    until = anchor + dt.timedelta(days=2)
    res.update(since=ts(since), until=ts(until), logging_stopped=logging_stopped)
    say(f"[{label}] review window {ts(since)} to {ts(until)} UTC")

    say(f"[{label}] checking known IOC paths, boot and cron hooks")
    res["ioc"] = check_known_iocs(t)
    res["hooks"] = check_hooks(t, ts(since), ts(until)) + check_httpd_conf(t, ts(since), ts(until))
    res["cores"] = list_cores(t)
    base_presence = {}
    if cand and cand["base"] != img_path:
        try:
            tb = Target.open(cand["base"])
            for r in res["ioc"]:
                if not r.get("present"):
                    continue
                if not exists(tb, r["path"]):
                    base_presence[r["path"]] = "no"
                elif r["type"] == "file":
                    same = entry_row(tb.fs.path(r["path"]))["sha256"] == r["sha256"]
                    base_presence[r["path"]] = "yes, same content" if same else "yes, different content"
                else:
                    base_presence[r["path"]] = "yes"
            res["base_checked"] = cand["base"]
        except Exception as ex:
            res["errors"].append(f"base disk check failed: {ex!r}")
    res["base_presence"] = base_presence

    say(f"[{label}] walking /flash and /var for hidden, implant-like and new items")
    res["discovery"] = discovery_walk(t, since, until)

    # log pass
    hunts = Hunts(args.ioc_table)
    inv = []
    xdir = tdir / "logs_xlsx"
    todo = [m for m in manifest if m["exported_to"]]
    say(f"[{label}] reading {len(todo)} exported log files"
        + ("" if args.skip_log_xlsx else " and converting each one to .xlsx"))
    t0 = time.time()
    for i, m in enumerate(todo, 1):
        local = Path(m["exported_to"])
        rel_local = local.relative_to(logs_dir)
        img_rel = m["image_path"][len("/var/log/"):] if m["group"] == "var_log" else m["image_path"]
        top = m["group"] == "var_log" and "/" not in img_rel
        out_x = xdir / rel_local.parent / (rel_local.name + ".xlsx")
        if not args.skip_log_xlsx:
            os.makedirs(lp(out_x.parent), exist_ok=True)
        r = process_log(local, img_rel, m["mtime"], top, out_x, hunts, args.skip_log_xlsx)
        inv.append({"image_path": m["image_path"], "exported_to": str(local), "xlsx": r["xlsx"], "kind": r["kind"],
                    "lines": r["lines"], "nul_bytes": r["nul_bytes"], "size": m["size"], "mtime": m["mtime"],
                    "btime": m["btime"], "sha256": m["sha256"], "note": r["note"]})
        if i % 100 == 0:
            say(f"[{label}]   {i}/{len(todo)} files ({time.time() - t0:.0f} s)")
    res["manifest"], res["inventory"], res["hunts"] = manifest, inv, hunts
    res["sections"] = summarize_sections(hunts)
    res["webshell"] = webshell_rows(hunts)
    res["webshell_by_ip"] = webshell_by_ip(res["webshell"])
    res["nsppe"] = nsppe_summary(hunts)
    res["ua_payloads"] = ua_payload_rows(hunts)
    res["shell_distinct"] = shell_distinct(hunts)
    clock_diff = sorted({(a, b) for a, b in hunts.clock if a != b})
    res["clock_note"] = ("syslog and GMT stamps agree in ns.log (appliance clock is UTC)" if hunts.clock and not clock_diff
                         else f"syslog and GMT stamps differ in ns.log, e.g. {clock_diff[:2]}" if clock_diff
                         else "no ns.log lines with both stamps")
    agent = next((m for m in manifest if m["image_path"] == "/var/tmp/.ux/agent.log" and m["exported_to"]), None)
    res["agent"] = agent_log_rows(agent["exported_to"], agent["mtime"]) if agent else []

    if args.foxit:
        say(f"[{label}] running Fox-IT iocitrix (the SUID check walks every file; this takes a few minutes)")
        try:
            res["foxit"] = run_foxit(t, args.foxit, tdir / "foxit_iocitrix_output.txt")
        except Exception as ex:
            res["foxit"] = None
            res["errors"].append(f"Fox-IT iocitrix failed: {ex!r}")
            say(f"[{label}]   Fox-IT iocitrix failed: {ex!r}")
    else:
        res["foxit"] = None

    vmsd = cand["vmsd"] if cand else []
    res["timeline"] = timeline(label, res["ioc"], res["hooks"], res["sections"], res["webshell_by_ip"], hunts,
                               res["agent"], logging_stopped, vmsd, ts(since), res["cores"])
    res["network"] = network_iocs(res, args)
    say(f"[{label}] writing the results workbook")
    res["workbook"] = str(tdir / f"{label}_Hunt_Results.xlsx")
    write_workbook(res, res["workbook"], args)
    res["status"] = "analyzed"
    return res


def network_iocs(res, args):
    rows = {}

    def add(ind, kind, role, where, first="", last="", n=1, detail=""):
        e = rows.setdefault((ind, role), {"indicator": ind, "type": kind, "role": role, "where": set(), "first": "",
                                          "last": "", "count": 0, "detail": detail})
        e["where"].add(where)
        e["count"] += n
        if first and (not e["first"] or first < e["first"]):
            e["first"] = first
        if last and last > e["last"]:
            e["last"] = last
    for rel, ln, t_utc, ip, _text in res["hunts"].hits.get("H", []):
        level, note = res["hunts"].ip_info(ip)
        add(norm_ip(ip) or ip, "ip", f"listed attacker IP ({level}) seen in logs",
            f"/var/log/{rel}" if not rel.startswith("/") else rel, t_utc, t_utc, detail=note)
    per_ip = collections.OrderedDict()
    for w in res["webshell_by_ip"]:
        e = per_ip.setdefault(w["ip"], {"n": 0, "body": 0, "paths": [], "first": "", "last": "", "status": []})
        e["n"] += w["count"]
        e["body"] += w["with_body"]
        e["paths"].append(w["path"])
        e["status"].append(w["status"])
        e["first"] = min(e["first"], w["first"]) if e["first"] else w["first"]
        e["last"] = max(e["last"], w["last"])
    for ip, e in per_ip.items():
        add(ip, "ip", "requested a webshell path", "; ".join(e["paths"]), e["first"], e["last"], e["n"],
            f"{e['body']} non-empty response(s); status codes {', '.join(e['status'])}")
    where_pb = "PITBOSS injected commands (hunts A and B)"
    for n in res["nsppe"]:
        for u in n["urls"].split():
            add(u, "url", "callback URL in an injected command", where_pb, n["first"], n["last"], n["count"])
            host = re.sub(r"^\w+://", "", u).split("/")[0].split("@")[-1].rsplit(":", 1)[0]
            add(host, "ip" if IPV4_RX.fullmatch(host) else "domain", "callback host in an injected command",
                where_pb, n["first"], n["last"], n["count"])
        for ip in n["ips_in_text"].split(", "):
            if ip:
                add(ip, "ip", "IP inside an injected command", where_pb, n["first"], n["last"], n["count"])
        for ip in n["client_ips"].split(", "):
            if ip:
                add(ip, "ip", "client IP of a forged pitboss / NSPPE login", where_pb, n["first"], n["last"], n["count"])
    for w in res.get("ua_payloads", []):
        where = f"/var/log/{w['file']}"
        if w["ip"]:
            add(norm_ip(w["ip"]) or w["ip"], "ip", "sent a base64 payload in the User-Agent (PITBOSS stage 1)", where,
                w["time"], w["time"], detail=f"payload: {w['kind']}")
        for u in w["urls"].split():
            add(u, "url", "URL inside a staged User-Agent payload", where, w["time"], w["time"])
            host = re.sub(r"^\w+://", "", u).split("/")[0].split("@")[-1].rsplit(":", 1)[0]
            add(host, "ip" if IPV4_RX.fullmatch(host) else "domain", "host inside a staged User-Agent payload", where,
                w["time"], w["time"])
        for ip in w["ips"].split(", "):
            if ip:
                add(ip, "ip", "IP inside a staged User-Agent payload", where, w["time"], w["time"])
    for m in res["manifest"]:
        if m["image_path"].endswith(".relay-ssh.log") and m["exported_to"]:
            txt = Path(m["exported_to"]).read_text(encoding="utf-8", errors="replace")
            for ip in sorted(set(IPV4_RX.findall(txt))):
                add(ip, "ip", "SSH server of the relay tunnel", m["image_path"], m["btime"], m["mtime"])
    for r in res["agent"]:
        if r["kind"] == "open" and r["dest_ip"]:
            add(f"{r['dest_ip']}:{r['dest_port']}", "ip:port", "destination reached through the tunnel (agent.log)",
                "/var/tmp/.ux/agent.log", f"{r['date']} {r['time']}", f"{r['date']} {r['time']}")
    out = []
    for e in rows.values():
        e["where"] = "; ".join(sorted(e["where"]))[:2000]
        out.append(e)
    out.sort(key=lambda x: (x["role"], -x["count"]))
    return out


def inventory_disk(t, label, tdir, args):
    """Non-NetScaler disk (for example the /var/crash disk): file listing with times."""
    rows, tops = [], collections.Counter()
    tops_bytes = collections.Counter()
    for e, st in walk(t, "/", limit=500000):
        if stat.S_ISDIR(st.st_mode):
            continue
        p = str(e)
        top = "/" + p.strip("/").split("/")[0] if "/" in p.strip("/") else "/"
        tops[top] += 1
        tops_bytes[top] += st.st_size
        rows.append([p, st.st_size, ts(st.st_mtime), ts(st.st_ctime), ts(getattr(st, "st_birthtime", None)),
                     oct(st.st_mode & 0o7777), st.st_uid])
    wbp = tdir / f"{label}_Disk_Inventory.xlsx"
    b = Book(wbp)
    b.table("Files", ["Path", "Size", "Modified (UTC)", "Changed (UTC)", "Created (UTC)", "Mode", "UID"], rows,
            [70, 14, 20, 20, 20, 8, 6])
    b.table("Top folders", ["Folder", "Files", "Bytes"],
            [[k, v, tops_bytes[k]] for k, v in sorted(tops.items())], [30, 10, 16])
    b.close()
    return {"inventory_rows": len(rows), "top_folders": [(k, v, tops_bytes[k]) for k, v in sorted(tops.items())],
            "workbook": str(wbp), "recent": sorted(rows, key=lambda r: r[2], reverse=True)[:25]}


# ------------------------------------------------------------------ results workbook
def write_workbook(res, path, args):
    b = Book(path)
    h = res["hunts"]
    info = res["info"]
    c = res["cand"] or {}
    present = [r for r in res["ioc"] if r.get("present")]
    flagged = [x for x in res["hooks"] if x["flags"]]
    leads = [d for d in res["discovery"] if d["category"].startswith(("IOC", "Lead"))]
    b.notes("README", [
        ("NetScaler hunt results: " + res["label"]),
        f"Image: {res['image']}",
        f"Generated {now_utc()} UTC by nshunt_engine {ENGINE_VERSION}. All times are UTC.",
        f"Review window: {res['since']} to {res['until']} (used by the Discovery sheet).",
        "Sheets: Summary, Target, Disk_Chain, IOC_Files (known implant paths with SHA-256), Hooks (boot script and "
        "crontab lines), Discovery (hidden, implant-like and new items), A to I (log hunt hits), Hunt_Summary, "
        "I_Shell_Distinct, Webshell_by_IP, Pitboss_Injections, Network_IOCs, Core_Dumps, Agent_Log (tunnel log), "
        "Logins_utx, Shell_Audit_Rare, Log_Integrity_NUL, Implant_Strings_by_File, FoxIT_iocitrix, Timeline, "
        "Log_Inventory, Export_Manifest.",
        "Web shells hidden in other files: Discovery reads the web folders (/var/netscaler/logon, /var/vpn, and static "
        "file types in /var/netscaler/gui) for PHP code in files that are not .php, checks the VPN / EPA client download "
        "folders for package-named files that are not real packages, and compares files in the web, download and temp "
        "folders with the SHA-256 values Unit 42 published.",
        "Persistence: /flash/nsconfig/httpd.conf and /flash/nsconfig/crontab are the persistent copies of /etc/httpd.conf "
        "and /etc/crontab (linked back into /etc at boot). Hooks lists the httpd.conf lines that switch the PHP engine "
        "on, alias a hidden file or a receiver.min.css name, or run a hidden file as PHP, and every crontab line.",
        "PITBOSS log poisoning (CVE-2026-88771): B lists forged pitboss / NSPPE records in any top-level /var/log file "
        "(a login user name that imitates a pitboss crash record, usually with a command after NSPPE); A lists Call Home "
        "crash checker lines (callhomedebug.log) that carry such a command, the step that runs it; B2 lists web requests "
        "whose User-Agent holds a base64 payload (the stage the command usually fetches and runs). Pitboss_Injections "
        "groups A and B by injected command. Genuine pitboss crash records are in C and D, never in B.",
        "Hooks: rc.netscaler, nsbefore.sh, nsafter.sh and crontabs are normal files; a line is a finding only when it "
        "has flags. A row with no line number is about the whole file (empty, or created or modified in the review "
        "window).",
        "Every exported log is also saved as its own .xlsx under logs_xlsx (same folder layout as the image).",
        "Line numbers refer to the exported log files under logs (decompressed for .gz).",
        "Hunt H searches every log for the built-in attacker IPs (the intrusion this tool was built on, plus public "
        f"reporting up to {REPORTED_AS_OF}) and any given with -IocIp. Level: confirmed, likely or suspected (may be a "
        "scanner) for the tool's own case; reported-infra (publicly reported C2, payload, staging or exfiltration "
        "server); reported-exploit (publicly reported exploit source or webshell user); reported-ioc (in a vendor's IOC "
        "table with no stated role); supplied (-IocIp). A hit is a "
        "lead: read the log line before calling it attacker activity.",
        "Syslog lines carry no year or zone: the year comes from the file's modified time. " + res.get("clock_note", ""),
        "agent.log lines carry only HH:MM:SS: dates are inferred backwards from the file's modified time; days with no "
        "lines cannot be detected, so treat the dates of older lines as approximate.",
        "Implant code (scripts, webshells) is recorded by SHA-256 and metadata only and is never exported.",
        "This workbook is generated automatically. Verify every finding against the source before reporting it.",
    ])
    summ = [["Hostname", info.get("hostname", "")], ["NetScaler version", info.get("os_version", "")],
            ["IPs", info.get("ips", "")], ["Install date", info.get("install_date", "")],
            ["Last activity (dissect)", info.get("last_activity", "")], ["Newest log modification", res["logging_stopped"]],
            ["Known IOC paths present", f"{len([r for r in present if r['query'] == r['path']])} of {len(KNOWN_IOC_PATHS)}"],
            ["Flagged hook lines", len(flagged)], ["Discovery IOCs and leads", len(leads)]]
    n_sup = sum(1 for lv, _n in h.ioc_ips.values() if lv == SUPPLIED_IOC[0])
    n_rep = sum(1 for lv, _n in h.ioc_ips.values() if lv.startswith("reported"))
    summ.append(["Attacker IPs searched (hunt H)",
                 f"{len(h.ioc_ips) - n_sup} built in ({len(h.ioc_ips) - n_sup - n_rep} from the tool's own case, {n_rep} "
                 f"from public reporting up to {REPORTED_AS_OF}) + {n_sup} supplied with -IocIp"])
    for s in res["sections"]:
        summ.append([f"Hunt {s['key']}: {s['title']}", s["count"]])
    summ += [["Distinct PITBOSS injected commands (A and B)",
              sum(1 for n in res["nsppe"] if n["injected"] != NO_COMMAND)],
             ["Staged User-Agent payloads (B2)", len(res["ua_payloads"])],
             ["Files matching a published SHA-256", sum(1 for d in res["discovery"]
                                                       if d["category"].startswith("IOC: SHA-256"))],
             ["Web files with PHP code that are not .php", sum(1 for d in res["discovery"] if "PHP code" in d["category"])],
             ["Core and crash dumps", len(res["cores"])],
             ["Tunnel opens in agent.log", sum(1 for r in res["agent"] if r["kind"] == "open")],
             ["Logs with NUL-byte runs", len({x[0] for x in h.nul})],
             ["Fox-IT iocitrix findings", "not run" if res["foxit"] is None else len(res["foxit"])],
             ["Exported files", sum(1 for m in res["manifest"] if m["exported_to"])],
             ["Logs converted to .xlsx", sum(1 for i in res["inventory"] if i["xlsx"])]]
    b.table("Summary", ["Item", "Value"], summ, [60, 60])
    b.table("Target", ["Field", "Value"], [[k, v] for k, v in info.items()] +
            [["Image opened", res["image"]], ["Image format", c.get("format", "")],
             ["Base disk checked", res.get("base_checked", "")],
             ["Review window", f"{res['since']} to {res['until']}"]], [24, 110])
    chain_rows = []
    segs = c.get("segments", [])
    fmt = c.get("format", "")
    for i, s in enumerate(segs, 1):
        own = "; ".join(p for p in c.get("problems", []) if s["file"] in p)
        chain_rows.append(["image file" if len(segs) == 1 else f"part {i} of {len(segs)}", s["file"], fmt, "", "", "",
                           "" if fmt.startswith("E01") else s["size"] // 512, s["size"], own or "ok"])
    if segs:
        for p in c.get("problems", []):
            if not any(s["file"] in p for s in segs):
                chain_rows.append(["problem", "", fmt, "", "", "", "", "", p])
    for d in c.get("chain_desc", []):
        chain_rows.append(["descriptor", d["file"], d["type"], d["cid"], d["parent_cid"], d["parent_hint"], "", "", ""])
    for e in c.get("extents", []):
        chain_rows.append(["extent", e["extent"], e["type"], "", "", e["descriptor"], e["sectors"], e["size"], e["status"]])
    for v in c.get("vmx", []):
        for dev, f in v["disks"]:
            found = any(Path(x).name.lower() == Path(f.replace("\\", "/")).name.lower() for x in c.get("chain", []))
            chain_rows.append(["vmx disk", f, dev, "", "", v["vmx"], "", "", "in this chain" if found else "not in this chain"])
    for s in c.get("vmsd", []):
        chain_rows.append(["snapshot", s["name"], "current" if s["current"] else "", "", "", "", "", "", s["created_utc"]])
    b.table("Disk_Chain", ["Kind", "File", "Type", "CID", "Parent CID", "Parent hint / descriptor", "Sectors", "Size",
                           "Status"], chain_rows, [12, 50, 14, 10, 10, 50, 12, 14, 30])
    bp = res.get("base_presence", {})
    b.table("IOC_Files", ["Path in image", "Present", "Type", "Size", "Mode", "UID", "Modified (UTC)", "Changed (UTC)",
                          "Created (UTC)", "SHA-256", "Symlink target", "Also in base disk", "Description", "Note"],
            ([r["path"], r.get("present"), r["type"], r["size"], r["mode"], r["uid"], r["mtime"], r["ctime"], r["btime"],
              r["sha256"], r["link"], bp.get(r["path"], ""), r.get("description", ""), r["note"]] for r in res["ioc"]),
            [55, 8, 8, 10, 8, 6, 20, 20, 20, 66, 20, 10, 60, 20])
    b.table("Hooks", ["File", "Line", "Text", "Flags", "File modified (UTC)", "File created (UTC)", "File SHA-256"],
            ([x["file"], x["line"], x["text"], x["flags"], x["mtime"], x["btime"], x["sha256"]] for x in res["hooks"]),
            [34, 6, 70, 60, 20, 20, 66])
    order = {"IOC": 0, "Lea": 1, "Con": 2}
    b.table("Discovery", ["Category", "Path", "Type", "Mode", "UID", "Size", "Modified (UTC)", "Changed (UTC)",
                          "Created (UTC)", "SHA-256 (web, download and temp folders)"],
            ([d["category"], d["path"], d["type"], d["mode"], d["uid"], d["size"], d["mtime"], d["ctime"], d["btime"],
              d.get("sha256", "")]
             for d in sorted(res["discovery"], key=lambda d: (order.get(d["category"][:3], 9), d["category"], d["path"]))),
            [50, 80, 8, 8, 6, 12, 20, 20, 20, 66])
    for key, sheet, title, pats, _rx in h.sections:
        rows = h.hits[key]
        if key == "E":
            b.table(sheet, ["Time (UTC)", "Client IP", "Method", "Path", "Status", "Bytes", "User agent", "File (/var/log)",
                            "Line"],
                    ([w["time"], w["ip"], w["method"], w["path"], w["status"], w["bytes"], w["ua"], w["file"], w["line"]]
                     for w in res["webshell"]), [20, 16, 8, 50, 8, 10, 50, 30, 8])
        elif key == "B2":
            b.table(sheet, ["Time (UTC)", "Client IP", "Method", "Path", "Status", "Payload marker", "Base64 length",
                            "Decodes to", "SHA-256 of the base64", "SHA-256 decoded", "Matches published hash",
                            "URLs inside", "IPs inside", "User agent (as logged)", "File (/var/log)", "Line"],
                    ([w["time"], w["ip"], w["method"], w["path"], w["status"], w["marker"], w["length"], w["kind"],
                      w["sha_b64"], w["sha_dec"], w["known"], w["urls"], w["ips"], w["ua"], w["file"], w["line"]]
                     for w in res["ua_payloads"]), [20, 16, 8, 40, 8, 22, 10, 18, 66, 66, 30, 50, 24, 80, 30, 8])
        elif key == "H":
            b.table(sheet, ["IP", "Level", "Known for", "Time (UTC)", "File (/var/log or image path)", "Line", "Text"],
                    ([r[3], *h.ip_info(r[3]), r[2], r[0], r[1], r[4]] for r in rows), [16, 10, 44, 20, 34, 8, 150])
        else:
            b.table(sheet, ["Time (UTC)", "File (/var/log)", "Line", "Text"], ([r[2], r[0], r[1], r[4]] for r in rows),
                    [20, 30, 8, 160])
    b.table("Hunt_Summary", ["Section", "Title", "Files read", "Hits", "Hits in /var/log top level", "First (UTC)",
                             "Last (UTC)", "Files (hits)"],
            ([s["key"], s["title"], s["files_filter"], s["count"], s["count_top_level"], s["first"], s["last"], s["files"]]
             for s in res["sections"]), [8, 50, 22, 8, 12, 20, 20, 80])
    b.table("I_Shell_Distinct", ["Command (section I, base64 runs shown as <B64>)", "Lines", "First (UTC)", "Last (UTC)",
                                 "Files (lines)"],
            ([c_["command"], c_["count"], c_["first"], c_["last"], c_["files"]] for c_ in res["shell_distinct"]),
            [120, 8, 20, 20, 60])
    b.table("Webshell_by_IP", ["Client IP", "Path", "Requests", "Non-empty responses", "First (UTC)", "Last (UTC)", "Days",
                               "Status codes"],
            ([w["ip"], w["path"], w["count"], w["with_body"], w["first"], w["last"], w["days"], w["status"]]
             for w in res["webshell_by_ip"]), [16, 50, 10, 12, 20, 20, 30, 30])
    b.table("Pitboss_Injections", ["Injected command (text after NSPPE, URL-decoded)", "Lines", "Forged records (B)",
                                   "Call Home lines (A)", "First (UTC)", "Last (UTC)", "Client IPs", "URLs", "IPs in text",
                                   "Files (lines)"],
            ([n["injected"], n["count"], n["forged"], n["callhome"], n["first"], n["last"], n["client_ips"], n["urls"],
              n["ips_in_text"], n["files"]] for n in res["nsppe"]), [80, 8, 10, 10, 20, 20, 24, 50, 20, 60])
    b.table("Network_IOCs", ["Indicator", "Type", "Role", "Seen in", "Count", "First (UTC)", "Last (UTC)", "Detail"],
            ([n["indicator"], n["type"], n["role"], n["where"], n["count"], n["first"], n["last"], n["detail"]]
             for n in res["network"]), [40, 8, 44, 60, 8, 20, 20, 50])
    b.table("Core_Dumps", ["Kind", "Path in image", "Size", "Modified (UTC)", "Created (UTC)"],
            ([c_["kind"], c_["path"], c_["size"], c_["mtime"], c_["btime"]] for c_ in res["cores"]), [28, 70, 14, 20, 20])
    b.table("Agent_Log", ["Date (inferred)", "Time", "Kind", "Destination IP", "Port", "Session", "Message"],
            ([r["date"], r["time"], r["kind"], r["dest_ip"], r["dest_port"], r["sid"], r["message"]] for r in res["agent"]),
            [14, 10, 12, 16, 8, 8, 90])
    b.table("Logins_utx", ["File (/var/log)", "Type", "Time (UTC)", "User", "Line (tty)", "Host", "PID"],
            ([rel, r["type"], r["time_utc"], r["user"], r["line"], r["host"], r["pid"]]
             for rel, r in sorted(h.utx, key=lambda x: (x[1]["time_utc"], x[0]))), [20, 16, 20, 14, 10, 30, 10])
    b.table("Shell_Audit_Rare", ["Command (sh_command, base64 runs shown as <B64>)", "Count", "First (UTC)", "Last (UTC)",
                                 "First file"],
            ([cmd, e[0], e[1], e[2], e[3]] for cmd, e in sorted(h.audit.items(), key=lambda x: (x[1][0], x[1][1]))),
            [120, 8, 20, 20, 24])
    b.table("Log_Integrity_NUL", ["File (/var/log or image path)", "Line", "NUL bytes", "Last time before (UTC)",
                                  "First time after (UTC)"], (list(x) for x in h.nul), [40, 8, 10, 20, 20])
    b.table("Implant_Strings_by_File", ["File (/var/log or image path)", "Lines", "First (UTC)", "Last (UTC)",
                                        "Strings matched"],
            ([f, e[0], e[1], e[2], ", ".join(sorted(e[3]))] for f, e in sorted(h.x.items(), key=lambda x: -x[1][0])),
            [44, 10, 20, 20, 60])
    if res["foxit"] is not None:
        b.table("FoxIT_iocitrix", ["Confidence", "Type", "Alert", "Path"],
                ([f["confidence"], f["type"], f["alert"], f["path"]] for f in res["foxit"]), [10, 20, 60, 80])
    b.table("Timeline", ["Time (UTC)", "Source", "Event", "Detail"],
            ([e["time"], e["source"], e["event"], e["detail"]] for e in res["timeline"]), [20, 30, 70, 70])
    b.table("Log_Inventory", ["Path in image", "Kind", "Lines or records", "NUL bytes", "Size", "Modified (UTC)",
                              "Created (UTC)", "SHA-256", "Exported to", "XLSX", "Note"],
            ([i["image_path"], i["kind"], i["lines"], i["nul_bytes"], i["size"], i["mtime"], i["btime"], i["sha256"],
              i["exported_to"], i["xlsx"], i["note"]] for i in res["inventory"]),
            [44, 22, 10, 10, 12, 20, 20, 66, 60, 60, 30])
    b.table("Export_Manifest", ["Group", "Path in image", "Exported to", "Size", "Mode", "UID", "Modified (UTC)",
                                "Changed (UTC)", "Created (UTC)", "SHA-256", "Note"],
            ([m["group"], m["image_path"], m["exported_to"], m["size"], m["mode"], m["uid"], m["mtime"], m["ctime"],
              m["btime"], m["sha256"], m["note"]] for m in res["manifest"]),
            [12, 44, 60, 12, 8, 6, 20, 20, 20, 66, 30])
    b.close()
    with open(lp(Path(path).parent / "logs" / "export_manifest.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["group", "image_path", "exported_to", "size", "mode", "uid", "mtime_utc", "ctime_utc", "btime_utc",
                    "sha256", "note"])
        for m in res["manifest"]:
            w.writerow([m["group"], m["image_path"], m["exported_to"], m["size"], m["mode"], m["uid"], m["mtime"],
                        m["ctime"], m["btime"], m["sha256"], m["note"]])


# ------------------------------------------------------------------ Word summary
def write_docx(path, run, results):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    def para(text, bold=False, italic=False, size=None, color=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor(*color)
        return p

    def bullet(text):
        doc.add_paragraph(text, style="List Bullet")

    def table(headers, rows, widths=None, mono=()):
        rows = list(rows)
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.style = "Table Grid"
        tb.autofit = False
        for i, h in enumerate(headers):
            cell = tb.rows[0].cells[i]
            cell.text = ""
            run_ = cell.paragraphs[0].add_run(h)
            run_.bold = True
            run_.font.size = Pt(9)
        for row in rows:
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run_ = cells[i].paragraphs[0].add_run(xl("" if v is None else str(v))[:1500])
                run_.font.size = Pt(8.5)
                if i in mono:
                    run_.font.name = "Consolas"
                    run_.font.size = Pt(7)
        if widths:
            for i, w in enumerate(widths):
                tb.columns[i].width = Inches(w)
                for cell in tb.columns[i].cells:
                    cell.width = Inches(w)
        doc.add_paragraph()
        return tb

    doc.add_heading("NetScaler disk image hunt: findings summary", 0)
    para(f"Evidence: {run['path']}")
    para(f"Generated {run['finished_utc']} UTC on {run['host']} by Invoke-NetScalerHunt / nshunt_engine {ENGINE_VERSION} "
         f"(dissect {run['versions'].get('dissect', '?')}, dissect.target {run['versions'].get('dissect.target', '?')}, "
         f"Python {run['python']}).")
    para("This summary is generated automatically from the disk images. Every statement points to a sheet of the "
         "per-appliance workbook; verify findings against the source before you report them. All times are UTC.",
         italic=True, color=(0x59, 0x59, 0x59))

    def version_str(info):
        k = (info.get("boot_kernel") or "").strip("/")
        fv = info.get("flash_version") or ""
        if k:
            return f"{k} (kernel booted)" + (f"; /flash/.version says {fv}" if fv else "")
        return info.get("os_version", "")

    doc.add_heading("Scope", 1)
    table(["Label", "Image opened", "Hostname", "Version", "IPs", "Newest log", "Status"],
          ([r["label"], r["image"] + (f" [{(r.get('cand') or {}).get('format')}]" if (r.get("cand") or {}).get("format") else ""),
            r.get("info", {}).get("hostname", ""), version_str(r.get("info", {})),
            r.get("info", {}).get("ips", "")[:200], r.get("logging_stopped", ""), r.get("status", "")] for r in results),
          [1.3, 2.6, 0.7, 1.6, 1.2, 1.3, 1.1])
    if run.get("notes"):
        for n in run["notes"]:
            bullet(n)
    iocs = run.get("ioc_ips", {})
    sup = iocs.get("supplied", [])
    sup_txt = ", ".join(sup[:20]) + (f" and {len(sup) - 20} more" if len(sup) > 20 else "")
    para(f"Hunt H searched every log for {len(iocs.get('built_in', []))} built-in attacker IPs ({len(CASE_IOC_IPS)} from "
         f"the intrusion this tool was built on, {len(iocs.get('built_in', [])) - len(CASE_IOC_IPS)} from public reporting "
         f"up to {REPORTED_AS_OF})" + (f" and {len(sup)} supplied with -IocIp ({sup_txt})" if sup else "")
         + ". The H_Attacker_IP_Lines sheet gives each IP's level (confirmed, likely, suspected, reported-infra, "
         "reported-exploit, reported-ioc or supplied) and what it was seen doing.", italic=True)

    hashes = collections.defaultdict(set)
    for r in results:
        for x in r.get("ioc", []):
            if x.get("sha256") and x.get("size"):
                hashes[(x["sha256"], Path(x["path"]).name)].add(r["label"])

    for r in results:
        if r.get("status") != "analyzed":
            doc.add_heading(f"{r['label']}: {r.get('status', '')}", 1)
            if r.get("top_folders"):
                para(f"File listing: {r.get('workbook', '')}")
                table(["Folder", "Files", "Bytes"], ([a, b_, f"{c_:,}"] for a, b_, c_ in r["top_folders"]), [3, 1, 2])
            for e in r.get("errors", []):
                bullet(e)
            continue
        info, h = r["info"], r["hunts"]
        secs = {s["key"]: s for s in r["sections"]}
        present = [x for x in r["ioc"] if x.get("present")]
        top_present = [x for x in present if x["query"] == x["path"]]
        doc.add_heading(f"{r['label']} ({info.get('hostname', '')}, {version_str(info)})", 1)
        para(f"Image: {r['image']}. Workbook: {r['workbook']}. Review window {r['since']} to {r['until']}.")

        doc.add_heading("At a glance", 2)
        found_any = False
        if top_present:
            found_any = True
            bullet(f"{len(top_present)} of {len(KNOWN_IOC_PATHS)} known IOC paths are present: "
                   + ", ".join(x["path"] for x in top_present) + ". (IOC_Files)")
        pb, pa, pu = secs.get("B", {}), secs.get("A", {}), secs.get("B2", {})
        if pb.get("count") or pa.get("count"):
            found_any = True
            cmds = [n for n in r["nsppe"] if n["injected"] != NO_COMMAND]
            text = (f"PITBOSS log poisoning (CVE-2026-88771): {pb.get('count', 0):,} forged pitboss / NSPPE log record(s)"
                    + (f", {pb['first']} to {pb['last']}" if pb.get("first") else "")
                    + (f", carrying {len(cmds)} distinct injected command(s)" if cmds
                       else ", none with a command after NSPPE (probes or tests)"))
            if pa.get("count"):
                text += (f"; {pa['count']:,} line(s) in the Call Home crash checker log (callhomedebug.log) carry an "
                         f"injected command, the step that runs it")
            else:
                text += ("; no callhomedebug.log line carries one (that log rotates, so this alone does not show the "
                         "commands never ran)")
            ips = sorted({ip for n in r["nsppe"] for ip in n["client_ips"].split(", ") if ip})
            if ips:
                text += f"; client IPs: {', '.join(ips[:15])}" + (f" and {len(ips) - 15} more" if len(ips) > 15 else "")
            bullet(text + ". (B_Pitboss_Log_Poisoning, A_CallHome_Exec, Pitboss_Injections)")
        if pu.get("count"):
            found_any = True
            ups = r["ua_payloads"]
            ua_ips = collections.Counter(w["ip"] for w in ups if w["ip"])
            kinds = collections.Counter(w["kind"] for w in ups)
            inside = sorted({u for w in ups for u in w["urls"].split()} | {i for w in ups for i in w["ips"].split(", ") if i})
            known = sorted({k for w in ups for k in w["known"].split("; ") if k})
            bullet(f"{pu['count']:,} web request(s) with a base64 payload in the User-Agent (PITBOSS stage 1: the injected "
                   f"command fetches it from the access log and runs it), {pu['first']} to {pu['last']}, from "
                   + ", ".join(f"{ip} ({n})" for ip, n in ua_ips.most_common(10))
                   + "; they decode to " + ", ".join(f"{k} x{v}" for k, v in kinds.most_common())
                   + (f"; URLs and IPs inside: {', '.join(inside[:12])}" if inside else "")
                   + (f"; {', '.join(known)}" if known else "") + ". (B2_Staged_UA_Payloads)")
        whole = {x["file"]: x for x in r["hooks"] if x["line"] == "" and x["flags"]}   # file-level rows
        by_file = collections.OrderedDict()
        for x in r["hooks"]:
            if x["line"] != "" and x["flags"]:
                by_file.setdefault(x["file"], []).append(x)
        for fname in list(by_file) + [f_ for f_ in whole if f_ not in by_file]:
            xs = by_file.get(fname, [])
            cmds = [x for x in xs if not x["text"].lstrip().startswith("#")]
            inst = [m.group(0) for x in xs for m in [re.search(r"installed from \S+ on .+$", x["flags"])] if m]
            flags = sorted({f_.strip() for x in cmds for f_ in x["flags"].split(";") if f_.strip()})
            w = whole.get(fname)
            wflags = [f_.strip() for f_ in w["flags"].split(";")] if w else []
            window = [f_ for f_ in wflags if "review window" in f_]
            if not cmds and not inst and not w:
                continue
            if cmds or inst or window:
                found_any = True
            parts = []
            if "empty crontab" in wflags:
                parts.append("exists but is empty")
            if window:
                parts.append(" and ".join(window) + f" (modified {w['mtime']}, created {w['btime'] or 'unknown'})")
            if cmds:
                parts.append("; ".join(f"line {x['line']} \"{x['text'].strip()}\"" for x in cmds))
            text = f"{fname}: " + "; ".join(parts) if parts else fname
            if inst:
                text += f" (crontab {inst[0]})"
            if flags:
                text += ". Flags: " + ", ".join(flags)
            if window and not cmds:
                text += ". Read the whole file and compare it with a known good copy"
            bullet(text + ". (Hooks)")
        if secs.get("E", {}).get("count"):
            found_any = True
            ips = len({w["ip"] for w in r["webshell_by_ip"]})
            body = sum(w["with_body"] for w in r["webshell_by_ip"])
            st_codes = collections.Counter(w["status"] for w in r["webshell"])
            bullet(f"{secs['E']['count']:,} requests to the webshell paths from {ips} client IP(s), {secs['E']['first']} to "
                   f"{secs['E']['last']} (status codes " + ", ".join(f"{k} x{v:,}" for k, v in st_codes.most_common(4))
                   + f"); {body:,} had a non-empty response body. (E_Webshell_Requests, Webshell_by_IP)")
        if secs.get("D", {}).get("count") or secs.get("C", {}).get("count"):
            bullet(f"{secs.get('D', {}).get('count', 0)} genuine packet engine (NSPPE) crash record(s) from pitboss and "
                   f"{secs.get('C', {}).get('count', 0)} nsaaad crash, restart-limit or pitboss reboot line(s); forged "
                   f"pitboss records are not counted here. (C, D)")
        cores = r.get("cores", [])
        if cores:
            ppe = sum(1 for c_ in cores if c_["kind"].startswith("packet engine"))
            inwin = sum(1 for c_ in cores if c_["mtime"] and c_["mtime"] >= r["since"])
            bullet(f"{len(cores)} core or crash dump(s) on disk ({ppe} from a packet engine), {inwin} written in the review "
                   f"window, newest {cores[-1]['mtime'] or 'unknown'}. Unit 42 says the CVE-2026-88772 DTLS exploit causes "
                   f"memory corruption or a crash, so line these up with the C and D records. (Core_Dumps)")
        if secs.get("F", {}).get("count"):
            found_any = True
            bullet(f"{secs['F']['count']} .relay installer line(s) in the shell audit log, {secs['F']['first']} to "
                   f"{secs['F']['last']}. (F_Relay_Installer)")
        if secs.get("G", {}).get("count"):
            found_any = True
            bullet(f"{secs['G']['count']:,} cron run(s) of implant scripts, {secs['G']['first']} to {secs['G']['last']}. "
                   f"(G_Implant_Cron_Runs)")
        if secs.get("H", {}).get("count"):
            found_any = True
            per_ip = collections.Counter(norm_ip(x[3]) or x[3] for x in h.hits["H"])
            bullet("Listed attacker IPs mentioned in logs: "
                   + ", ".join(f"{ip} ({n:,} line(s), {h.ip_info(ip)[0]})" for ip, n in per_ip.most_common())
                   + ". Read the lines before calling them attacker activity. (H_Attacker_IP_Lines)")
        opens = [a for a in r["agent"] if a["kind"] == "open"]
        if opens:
            found_any = True
            dests = collections.Counter(f"{a['dest_ip']}:{a['dest_port']}" for a in opens)
            bullet(f"Tunnel log /var/tmp/.ux/agent.log records {len(opens)} connection(s) opened to: "
                   + ", ".join(f"{d} ({n})" for d, n in dests.most_common()) + ". Dates are inferred. (Agent_Log)")
        nul_files = collections.Counter(x[0] for x in h.nul)
        if nul_files:
            bullet(f"{len(nul_files)} log file(s) contain runs of NUL bytes, which can mean the file was truncated while "
                   f"the logger kept writing (or an unclean write): " + ", ".join(nul_files) + ". (Log_Integrity_NUL)")
        if r["foxit"] is not None:
            types = collections.Counter(f["type"] for f in r["foxit"])
            bullet(f"Fox-IT iocitrix: {len(r['foxit'])} finding(s)"
                   + (": " + ", ".join(f"{k} {v}" for k, v in types.most_common()) if types else "")
                   + ". file/timestomp hits are often caused by upgrades. (FoxIT_iocitrix)")
        leads = [d for d in r["discovery"] if d["category"].startswith("Lead")]
        hash_hits = [d for d in r["discovery"] if d["category"].startswith("IOC: SHA-256")]
        if hash_hits:
            found_any = True
            bullet("Files matching a published hash: " + "; ".join(
                f"{d['path']} ({d['category'][len('IOC: SHA-256 matches '):]})" for d in hash_hits[:10]) + ". (Discovery)")
        pkg = [d for d in leads if "client download folder" in d["category"]]
        if pkg:
            found_any = True
            bullet("VPN / EPA client download folder: " + "; ".join(f"{d['path']} ({d['category'].split(': ', 2)[-1]})"
                                                                  for d in pkg[:10])
                   + ". Unit 42 reported the actor hosting webshells there under package names. (Discovery)")
        webphp = [d for d in leads if d["category"].startswith("Lead: web folder: PHP code")]
        if webphp:
            found_any = True
            bullet("PHP code inside non-PHP files in the web folders (how this actor hid its web shells, reached through "
                   "aliases such as receiver.min.css): " + ", ".join(d["path"] for d in webphp[:10])
                   + (f" and {len(webphp) - 10} more" if len(webphp) > 10 else "") + ". (Discovery)")
        drop = [d for d in leads if d["path"] == LOGON_DROP_DIR]
        if drop:
            found_any = True
            bullet(f"{LOGON_DROP_DIR} was created in the review window ({drop[0]['btime']}); Unit 42 reports the actor "
                   f"creating it as the drop folder for its web shell. (Discovery)")
        if leads:
            bullet(f"{len(leads)} item(s) on disk to review (hidden, new in /flash, /var/tmp or /var/nstmp, written by the "
                   f"web server in the window, PHP code in a non-PHP web file, or a suspect file in the VPN / EPA client "
                   f"download folders). (Discovery)")
        if not found_any:
            bullet("None of the known implant paths, hooks or log indicators were found in this image.")

        if present:
            doc.add_heading("Implant and persistence files (known IOC paths)", 2)
            table(["Path in image", "Type", "Size", "Created (UTC)", "Modified (UTC)", "SHA-256", "In base disk"],
                  ([x["path"], x["type"], x["size"], x["btime"], x["mtime"], x["sha256"],
                    r["base_presence"].get(x["path"], "")] for x in present),
                  [2.6, 0.45, 0.55, 1.15, 1.15, 3.3, 0.6], mono=(5,))
        if r["hooks"]:
            flagged_h = [x for x in r["hooks"] if x["flags"]]
            quiet = [x for x in r["hooks"] if not x["flags"] and x["line"] != ""]
            quiet_files = sorted({x["file"] for x in quiet})
            doc.add_heading("Boot script and crontab lines with flags", 2)
            if quiet:
                para(f"{len(quiet)} other line(s) in {', '.join(quiet_files)} have no flags and are not findings (these "
                     f"files hold ordinary settings on most appliances); all lines are on the Hooks sheet.", italic=True)
            if flagged_h:
                table(["File", "Line", "Text", "Flags"],
                      ([x["file"], x["line"], x["text"], x["flags"]] for x in flagged_h), [2.2, 0.4, 3.6, 3.6])
            else:
                para("No boot script or crontab line was flagged.")
        doc.add_heading("Log hunts", 2)
        table(["Section", "What it looks for", "Hits", "First (UTC)", "Last (UTC)", "Files (hits)"],
              ([s["key"], s["title"], f"{s['count']:,}", s["first"], s["last"], s["files"][:600]] for s in r["sections"]),
              [0.5, 2.6, 0.6, 1.3, 1.3, 3.5])
        if r["shell_distinct"]:
            doc.add_heading("Shell history commands that match section I (rarest first)", 2)
            para(f"{len(r['shell_distinct'])} distinct commands in {secs.get('I', {}).get('count', 0):,} lines; the 30 rarest "
                 f"are shown. Frequent ones are usually scripts that run on a schedule. (I_Shell_Distinct)", italic=True)
            table(["Command", "Lines", "First (UTC)", "Last (UTC)"],
                  ([c_["command"], c_["count"], c_["first"], c_["last"]] for c_ in r["shell_distinct"][:30]),
                  [5.6, 0.6, 1.4, 1.4])
        if r["webshell_by_ip"]:
            doc.add_heading("Webshell requests by client IP", 2)
            table(["Client IP", "Path", "Requests", "Non-empty", "First (UTC)", "Last (UTC)", "Status codes"],
                  ([w["ip"], w["path"], f"{w['count']:,}", w["with_body"], w["first"], w["last"], w["status"]]
                   for w in r["webshell_by_ip"][:25]), [1.2, 2.8, 0.7, 0.7, 1.3, 1.3, 1.6])
        if r["nsppe"]:
            doc.add_heading("PITBOSS log poisoning: injected commands (CVE-2026-88771)", 2)
            para("Each forged record is a login whose user name imitates a pitboss crash record; the Call Home crash "
                 "checker runs the text after NSPPE on an unpatched appliance. Forged = log records (B); Call Home = "
                 "callhomedebug.log lines (A). The full lines are on the B and A sheets.", italic=True)
            table(["Injected command", "Forged", "Call Home", "First (UTC)", "Last (UTC)", "Client IPs"],
                  ([n["injected"], n["forged"], n["callhome"], n["first"], n["last"], n["client_ips"]]
                   for n in r["nsppe"][:60]), [4.6, 0.55, 0.65, 1.3, 1.3, 1.4])
        if r["ua_payloads"]:
            doc.add_heading("PITBOSS stage 1: base64 payloads staged in the User-Agent", 2)
            para("The payload content is not reproduced here; the B2 sheet has its SHA-256, what it decodes to and the "
                 "URLs and IPs inside it. Decode the User-Agent from the log line to read it.", italic=True)
            table(["Time (UTC)", "Client IP", "Path", "Status", "Decodes to", "URLs and IPs inside"],
                  ([w["time"], w["ip"], w["path"], w["status"], w["kind"],
                    " ".join(x for x in (w["urls"], w["ips"]) if x)] for w in r["ua_payloads"][:40]),
                  [1.3, 1.2, 2.4, 0.6, 1.2, 3.1])
        if r["network"]:
            doc.add_heading("Network indicators", 2)
            table(["Indicator", "Role", "Count", "First (UTC)", "Last (UTC)", "Detail"],
                  ([n["indicator"], n["role"], n["count"], n["first"], n["last"], n["detail"]] for n in r["network"][:60]),
                  [2.2, 2.4, 0.6, 1.3, 1.3, 2.0])
        logins = [(rel, x) for rel, x in h.utx if x["type"] == "USER_PROCESS" and x["time_utc"] >= r["since"]]
        if logins:
            doc.add_heading("Logins recorded in utx.log (review window)", 2)
            agg = collections.OrderedDict()
            for rel, x in sorted(logins, key=lambda y: y[1]["time_utc"]):
                e = agg.setdefault((x["user"], x["host"]), [0, x["time_utc"], x["time_utc"]])
                e[0] += 1
                e[2] = x["time_utc"]
            table(["User", "From host", "Logins", "First (UTC)", "Last (UTC)"],
                  ([u, hst, n, f, l] for (u, hst), (n, f, l) in agg.items()), [1.5, 2.5, 0.8, 1.6, 1.6])
        if nul_files:
            doc.add_heading("Log integrity: NUL-byte runs", 2)
            table(["File", "Line", "NUL bytes", "Last time before", "First time after"], (list(x) for x in h.nul[:40]),
                  [2.6, 0.6, 1.0, 1.6, 1.6])
        if leads or hash_hits:
            doc.add_heading("Items on disk to review", 2)
            table(["Category", "Path", "Size", "UID", "Created (UTC)", "Modified (UTC)"],
                  ([d["category"], d["path"], d["size"], d["uid"], d["btime"], d["mtime"]] for d in (hash_hits + leads)[:60]),
                  [2.4, 3.6, 0.7, 0.5, 1.3, 1.3])
        cores = r.get("cores", [])
        if cores:
            doc.add_heading("Core and crash dumps", 2)
            para("Process, packet engine and kernel dumps found on the disk. They are not findings by themselves; line "
                 "their times up with the C and D crash records. (Core_Dumps)", italic=True)
            table(["Kind", "Path", "Size", "Modified (UTC)", "Created (UTC)"],
                  ([c_["kind"], c_["path"], f"{c_['size']:,}", c_["mtime"], c_["btime"]] for c_ in cores[-40:]),
                  [1.8, 3.8, 0.9, 1.3, 1.3])
        if r["foxit"]:
            doc.add_heading("Fox-IT iocitrix findings (excluding file/timestomp)", 2)
            nts = [f for f in r["foxit"] if f["type"] != "file/timestomp"]
            if nts:
                table(["Confidence", "Type", "Alert", "Path"], ([f["confidence"], f["type"], f["alert"], f["path"]] for f in nts),
                      [0.8, 1.3, 3.0, 4.0])
            else:
                para("Only file/timestomp findings; see the FoxIT_iocitrix sheet.")
        doc.add_heading("Timeline (UTC, within the review window)", 2)
        tl = [e for e in r["timeline"] if e["time"] >= r["since"]]
        if len(tl) > 120:
            para(f"{len(tl)} events; the first 120 are shown. The full list is on the Timeline sheet.", italic=True)
        table(["Time (UTC)", "Source", "Event", "Detail"], ([e["time"], e["source"], e["event"], e["detail"]] for e in tl[:120]),
              [1.3, 1.8, 3.5, 2.8])

    shared = [(k, v) for k, v in hashes.items() if len(v) > 1]
    if shared:
        doc.add_heading("Identical files across images", 1)
        table(["File name", "SHA-256", "Images"], ([k[1], k[0], ", ".join(sorted(v))] for k, v in sorted(shared)),
              [2.0, 4.4, 3.2], mono=(1,))

    doc.add_heading("Caveats and gaps", 1)
    for n in run.get("caveats", []):
        bullet(n)
    for r in results:
        cand = r.get("cand") or {}
        for p in cand.get("problems", []):
            if cand.get("chain_desc"):
                bullet(f"{r['label']}: disk file {p}. Data stored in the missing part cannot be read.")
            else:
                bullet(f"{r['label']}: image problem: {p}.")
        for e in r.get("errors", []):
            bullet(f"{r['label']}: {e}")
        if r.get("clock_note"):
            bullet(f"{r['label']}: {r['clock_note']}.")
    doc.add_heading("Output files", 1)
    for r in results:
        if r.get("workbook"):
            bullet(f"{r['label']}: {r['workbook']}")
    bullet("Every exported log as .xlsx: <label>\\logs_xlsx. Raw exported logs: <label>\\logs (export_manifest.csv "
           "has the SHA-256 and original UTC times of each file).")
    bullet("SHA256SUMS.txt in the output folder lists the SHA-256 of every file this run wrote, except engine.log and "
           "the PowerShell transcript, which are still being written when the hashes are taken.")
    doc.save(lp(path))


# ------------------------------------------------------------------ main
def main():
    global LOG
    ap = argparse.ArgumentParser(description="NetScaler disk image hunt (read-only).")
    ap.add_argument("--path", required=True,
                    help="folder to search for disk images (VMDK, .dd/.raw/.img, split .001, E01), or one image file")
    ap.add_argument("--out", required=True, help="output folder (must be outside the evidence folder)")
    ap.add_argument("--since", help="start of the review window, YYYY-MM-DD")
    ap.add_argument("--ioc-ip", nargs="*", default=[], metavar="IP",
                    help="extra attacker IPs (IPv4 or IPv6) to search every log for; added to the built-in list")
    ap.add_argument("--foxit", help="path to Fox-IT iocitrix.py")
    ap.add_argument("--skip-log-xlsx", action="store_true")
    ap.add_argument("--include-base", action="store_true")
    args = ap.parse_args()
    extra, bad = [], []
    for x in args.ioc_ip or []:
        for part in re.split(r"[,;\s]+", x.strip()):
            if part:
                ip = norm_ip(part)
                if ip is None:
                    bad.append(part)
                else:
                    extra.append(ip)
    if bad:
        ap.error("--ioc-ip: not an IP address: " + ", ".join(bad))
    builtin = [norm_ip(ip) for ip, _lv, _note in BUILTIN_IOC_IPS]
    extra = [ip for ip in dict.fromkeys(extra) if ip not in builtin]
    src, out = Path(args.path).resolve(), Path(args.out).resolve()
    if not src.exists():
        print(f"ERROR: {src} does not exist.", file=sys.stderr)
        return 2
    evidence = src.parent if src.is_file() else src
    if out == evidence or evidence in out.parents:
        print("ERROR: the output folder must not be inside the evidence folder.", file=sys.stderr)
        return 3
    os.makedirs(lp(out / "_tool"), exist_ok=True)
    LOG = open(lp(out / "_tool" / "engine.log"), "a", encoding="utf-8")
    logging.basicConfig(filename=lp(out / "_tool" / "dissect_messages.log"), level=logging.WARNING,
                        format="%(asctime)s %(name)s %(levelname)s %(message)s")
    run = {"path": str(src), "out": str(out), "started_utc": now_utc(), "host": platform.node(),
           "python": platform.python_version(), "engine": ENGINE_VERSION,
           "versions": {p: pkg_version(p) for p in ("dissect", "dissect.target", "XlsxWriter", "python-docx", "tabulate")},
           "args": dict(vars(args)), "ioc_ips": {"built_in": builtin, "supplied": extra}, "notes": [], "caveats": []}
    args.ioc_table = ioc_ip_table(extra)
    say(f"nshunt_engine {ENGINE_VERSION}: evidence {src}, output {out}")
    say(f"attacker IPs searched (hunt H): {len(builtin)} built in ({len(CASE_IOC_IPS)} from the tool's own case, "
        f"{len(builtin) - len(CASE_IOC_IPS)} from public reporting up to {REPORTED_AS_OF})"
        + (f" + {len(extra)} supplied: {', '.join(extra[:20])}" + (f" and {len(extra) - 20} more" if len(extra) > 20 else "")
           if extra else ""))
    cands, notes = discover(src)
    run["notes"] = notes
    for n in notes:
        say(f"note: {n}")
    if not cands:
        say("No supported disk images were found (VMDK, .dd/.raw/.img, split .001, E01).")
        json.dump(run, open(lp(out / "_tool" / "run_info.json"), "w"), indent=2)
        return 2
    for c in cands:
        if c.get("chain_desc"):
            what = f"disk chain {c['label']}: {' -> '.join(Path(x).name for x in c['chain'])}"
        else:
            what = f"{c['format']} image {c['label']}: {', '.join(s['file'] for s in c['segments']) or Path(c['leaf']).name}"
        say(f"found {what}" + (f"  [{'; '.join(c['problems'])}]" if c["problems"] else ""))
    results = []
    for c in cands:
        jobs = [(c["leaf"], c["label"])]
        if args.include_base and c["base"] != c["leaf"]:
            jobs.append((c["base"], c["label"] + "_base"))
        for img, label in jobs:
            try:
                results.append(hunt_image(c if img == c["leaf"] else dict(c, base=img), img, label, out, args, run))
            except Exception as ex:
                say(f"[{label}] FAILED: {ex!r}")
                LOG.write(traceback.format_exc())
                results.append({"label": label, "image": img, "cand": c, "status": f"error: {ex!r}", "errors": [repr(ex)]})
    run["caveats"] = [
        "RAM-only paths (/tmp, /etc and the running configuration) are not on the disk images and were not examined.",
        "Implant scripts and webshells are recorded by SHA-256 and metadata only; they are not exported, so antivirus does "
        "not react to the output.",
        "Syslog lines carry no year or time zone; the year comes from each file's modified time and the appliance clock is "
        "assumed to be UTC (see the per-image clock check below).",
        "agent.log lines carry only HH:MM:SS; their dates are inferred backwards from the file's modified time.",
        "The Discovery sheet uses file creation (birth) times; files copied or restored with old times can be missed.",
        "Fox-IT file/timestomp findings are frequent after upgrades and are leads, not proof.",
    ]
    run["finished_utc"] = now_utc()
    analyzed = [r for r in results if r.get("status") == "analyzed"]
    say("writing the Word summary")
    docx_path = out / "NetScaler_Findings_Summary.docx"
    try:
        write_docx(docx_path, run, results)
        run["docx"] = str(docx_path)
    except Exception as ex:
        say(f"Word summary failed: {ex!r}")
        LOG.write(traceback.format_exc())
    run["results"] = [{"label": r["label"], "image": r["image"], "format": (r.get("cand") or {}).get("format", ""),
                       "status": r.get("status"), "workbook": r.get("workbook", ""),
                       "sections": {s["key"]: s["count"] for s in r.get("sections", [])}} for r in results]
    json.dump(run, open(lp(out / "_tool" / "run_info.json"), "w"), indent=2, default=str)
    say("hashing the output files of this run")
    LOG.flush()
    lines = []
    # Only what this run wrote: one folder per image, _tool and the summary. Other content of --out is left alone.
    roots = [out / r["label"] for r in results] + [out / "_tool", docx_path]
    skip = {"SHA256SUMS.txt", "_tool/engine.log", "_tool/powershell_transcript.txt"}
    for root_ in dict.fromkeys(roots):
        if root_.is_file():
            walked = [(lp(root_.parent), [], [root_.name])]
        elif root_.is_dir():
            walked = os.walk(lp(root_))
        else:
            continue
        for dirpath, _d, files in walked:
            for fn in sorted(files):
                full = Path(dirpath) / fn
                plain = str(full)[4:] if str(full).startswith("\\\\?\\") else str(full)   # drop the long-path prefix
                rel = os.path.relpath(plain, str(out)).replace("\\", "/")
                if rel in skip:
                    continue
                lines.append(f"{sha256_file(full)}  {rel}")
    with open(lp(out / "SHA256SUMS.txt"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(sorted(lines, key=lambda x: x[66:])) + "\n")
    for r in results:
        say(f"{r['label']}: {r.get('status')}" + (f" -> {r.get('workbook')}" if r.get("workbook") else ""))
    say(f"done. Summary: {docx_path}")
    return 0 if analyzed else 1


if __name__ == "__main__":
    sys.exit(main())
